"""Professional multi-campaign Excel workbook export.

Presentation layer only: every number comes from the shared reporting
services (apps.reports.services), so the workbook always agrees with the
Campaign Report UI, the CSV export and the simple XLSX export.

Layout: a dashboard-style Summary (KPI cards, rate strips, campaign
comparison table, charts, navigation) plus one interactive Excel-Table
data sheet per enabled section.
"""
from io import BytesIO

from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table, TableStyleInfo

from .services import (
    LIFECYCLE_LABELS,
    build_report_filename,
    get_click_activity_rows,
    get_recipient_lifecycle_rows,
    get_reminder_activity_rows,
    reminder_status_label,
)

# ---------------------------------------------------------------------------
# Restrained visual identity (slate/navy family already used by the app).
# ---------------------------------------------------------------------------
NAVY_FILL = PatternFill('solid', fgColor='0F172A')
HEADER_FILL = PatternFill('solid', fgColor='1E293B')
SECTION_FILL = PatternFill('solid', fgColor='E2E8F0')
CARD_FILL = PatternFill('solid', fgColor='F8FAFC')
STRIP_FILL = PatternFill('solid', fgColor='F1F5F9')
NAV_FILL = PatternFill('solid', fgColor='E8F0FE')

TITLE_FONT = Font(name='Calibri', size=18, bold=True, color='FFFFFF')
SUBTITLE_FONT = Font(name='Calibri', size=10, italic=True, color='CBD5E1')
SECTION_FONT = Font(name='Calibri', size=11, bold=True, color='0F172A')
HEADER_FONT = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
BODY_FONT = Font(name='Calibri', size=10, color='0F172A')
SMALL_FONT = Font(name='Calibri', size=9, color='475569')
NOTE_FONT = Font(name='Calibri', size=10, italic=True, color='64748B')
META_LABEL_FONT = Font(name='Calibri', size=10, bold=True, color='475569')
META_VALUE_FONT = Font(name='Calibri', size=10, bold=True, color='0F172A')
CARD_LABEL_FONT = Font(name='Calibri', size=9, bold=True, color='64748B')
CARD_VALUE_FONT = Font(name='Calibri', size=20, bold=True, color='0F172A')
CARD_SUB_FONT = Font(name='Calibri', size=8, color='64748B')
STRIP_LABEL_FONT = Font(name='Calibri', size=9, bold=True, color='475569')
STRIP_VALUE_FONT = Font(name='Calibri', size=14, bold=True, color='0F172A')
LINK_FONT = Font(name='Calibri', size=10, bold=True, color='1D4ED8',
                 underline='single')
TOTAL_FONT = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
SHEET_TITLE_FONT = Font(name='Calibri', size=11, bold=True, color='475569')

THIN_BORDER = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0'),
)
CARD_EDGE = Side(style='medium', color='0F172A')
HEADER_BORDER = Border(bottom=Side(style='medium', color='0F172A'))

CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT_WRAP = Alignment(horizontal='left', vertical='center', wrap_text=True)
RIGHT_ALIGN = Alignment(horizontal='right', vertical='center')
LEFT_MID = Alignment(horizontal='left', vertical='center')

COUNT_FORMAT = '#,##0'
PERCENT_FORMAT = '0.0%'
DATETIME_FORMAT = 'DD MMM YYYY HH:MM'
TABLE_STYLE_NAME = 'TableStyleMedium2'

# Chart series palette: slate + one restrained blue + lifecycle accents.
CHART_SENT = '64748B'
CHART_DELIVERED = '2563EB'
CHART_COMPLETED = '059669'
CHART_INCOMPLETE = 'D97706'

# Conditional-formatting treatments: always text + subtle fill, never
# color alone. (fill, font color)
CF_STYLES = {
    'COMPLETED': ('E6F4EA', '137333'),
    'INCOMPLETE': ('FFF4DE', '92400E'),
    'ELIGIBLE': ('E8F0FE', '1A56DB'),
    'SUPPRESSED': ('F1F5F4', '64748B'),
    'YES': (None, '137333'),
    'Sent': ('E6F4EA', '137333'),
    'Failed': ('FDE8E8', '991B1B'),
    'Suspected Bot': ('FFF4DE', '92400E'),
}

EMPTY_NOTE = 'No records match the selected filters.'


def excel_dt(value):
    """Convert an aware datetime to a naive local datetime for openpyxl."""
    if value is None:
        return None
    from django.utils import timezone
    if timezone.is_aware(value):
        value = timezone.localtime(value)
        value = value.replace(tzinfo=None)
    return value


def _section_on(sections, key):
    if key not in sections:
        return True
    raw = sections[key]
    if isinstance(raw, str):
        return raw.strip().lower() not in ('false', '0', 'no', 'off', '')
    return bool(raw)


def _internal_link(cell, sheet, target='A1', display=None):
    """Turn a cell into an internal navigation hyperlink."""
    cell.hyperlink = Hyperlink(
        ref=cell.coordinate, location="'%s'!%s" % (sheet, target),
        display=display if display is not None else cell.value)


def _summary_title(ws, ncols, campaigns_label, generated_label,
                   period_label, filters_label):
    """Navy title band + subtitle + metadata grid. Returns next free row."""
    for row, value, font, height in (
            (1, 'CAMPAIGN REPORT', TITLE_FONT, 32),
            (2, 'Executive Performance & Recipient Lifecycle Overview',
             SUBTITLE_FONT, 20)):
        ws.merge_cells(start_row=row, start_column=1,
                       end_row=row, end_column=ncols)
        cell = ws.cell(row=row, column=1, value=value)
        cell.font = font
        cell.alignment = Alignment(horizontal='left', vertical='center',
                                    indent=1)
        ws.row_dimensions[row].height = height
        for col in range(1, ncols + 1):
            ws.cell(row=row, column=col).fill = NAVY_FILL
    # Metadata: column A is a narrow margin, content lives in B..I.
    ws.cell(row=3, column=2, value='Campaign(s):').font = META_LABEL_FONT
    ws.merge_cells(start_row=3, start_column=3, end_row=3, end_column=ncols - 1)
    ws.cell(row=3, column=3, value=campaigns_label).font = META_VALUE_FONT
    ws.row_dimensions[3].height = 18
    ws.cell(row=4, column=2, value='Reporting period:').font = META_LABEL_FONT
    ws.merge_cells(start_row=4, start_column=4, end_row=4, end_column=5)
    ws.cell(row=4, column=4, value=period_label).font = META_VALUE_FONT
    ws.cell(row=4, column=6, value='Generated on:').font = META_LABEL_FONT
    ws.merge_cells(start_row=4, start_column=7, end_row=4,
                   end_column=ncols - 1)
    ws.cell(row=4, column=7, value=generated_label).font = META_VALUE_FONT
    ws.row_dimensions[4].height = 18
    ws.cell(row=5, column=2, value='Applied filters:').font = META_LABEL_FONT
    ws.merge_cells(start_row=5, start_column=4, end_row=5,
                   end_column=ncols - 1)
    filters_cell = ws.cell(row=5, column=4, value=filters_label)
    filters_cell.font = BODY_FONT
    filters_cell.alignment = LEFT_WRAP
    ws.row_dimensions[5].height = 30
    for row in (3, 4, 5):
        ws.cell(row=row, column=2).alignment = LEFT_MID
        ws.cell(row=row, column=4).alignment = LEFT_WRAP
    return 7  # row 6 stays blank


def _section_heading(ws, row, ncols, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = SECTION_FONT
    cell.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[row].height = 22
    for col in range(1, ncols + 1):
        ws.cell(row=row, column=col).fill = SECTION_FILL
    return row + 1


def _header_row(ws, row, headers, widths=None):
    for idx, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=idx, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = HEADER_BORDER
    ws.row_dimensions[row].height = 28
    if widths:
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width


def _add_table(ws, name, ref):
    """Attach a banded, filterable Excel Table."""
    table = Table(
        displayName=name, ref=ref,
        tableStyleInfo=TableStyleInfo(
            name=TABLE_STYLE_NAME, showFirstColumn=False,
            showLastColumn=False, showRowStripes=True,
            showColumnStripes=False))
    table.autoFilter = AutoFilter(ref=ref)
    ws.add_table(table)
    return table


def _finish_sheet_common(ws, title):
    """Landscape print setup shared by every sheet."""
    from openpyxl.worksheet.properties import PageSetupProperties
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.oddFooter.left.text = 'Campaign Report — %s' % title
    ws.oddFooter.left.font = 'Calibri,Regular'
    ws.oddFooter.left.size = 9
    ws.oddFooter.center.text = 'Page &P of &N'
    ws.oddFooter.center.font = 'Calibri,Regular'
    ws.oddFooter.center.size = 9
    ws.sheet_view.showGridLines = False


def _cf_equal(ws, ref, text):
    """Conditional format: cells equal to text get the subtle CF_STYLES look."""
    fill_hex, font_color = CF_STYLES[text]
    rule = CellIsRule(
        operator='equal', formula=['"%s"' % text],
        font=Font(name='Calibri', size=10, bold=True, color=font_color),
        fill=(PatternFill('solid', fgColor=fill_hex)
              if fill_hex is not None else None))
    ws.conditional_formatting.add(ref, rule)


def _aggregate(rows, sent=None, delivered=None):
    """Aggregates over one filtered row set.

    sent/delivered count INITIAL/REMINDER messages with the Campaign
    Report's exact status sets (TEST messages excluded); opened/clicked
    count recipients, exactly like the report's unique counts. Rate
    bases match the Campaign Report: open/click rates over delivered,
    completion rate over recipients.
    """
    recipients = len(rows)
    if sent is None:
        sent = sum(1 for r in rows if r['has_sent'])
    if delivered is None:
        delivered = sum(1 for r in rows if r['has_delivered'])
    opened = sum(1 for r in rows if r['has_opened'])
    clicked = sum(1 for r in rows if r['has_clicked'])
    completed = sum(1 for r in rows if r['contact_status'] == 'USED')
    incomplete = recipients - completed
    eligible = sum(1 for r in rows if r['reminder_eligible'])
    suppressed = recipients - eligible
    return {
        'recipients': recipients,
        'sent': sent,
        'delivered': delivered,
        'opened': opened,
        'clicked': clicked,
        'completed': completed,
        'incomplete': incomplete,
        'eligible': eligible,
        'suppressed': suppressed,
        'open_rate': round(opened / delivered * 100, 1) if delivered else 0.0,
        'click_rate': round(clicked / delivered * 100, 1) if delivered else 0.0,
        'completion_rate': round(completed / recipients * 100, 1) if recipients else 0.0,
    }


def _filters_summary(filters):
    filters = filters or {}
    lifecycle = str(filters.get('lifecycle', 'all') or 'all').lower()
    parts = [LIFECYCLE_LABELS.get(lifecycle, 'All recipients')]
    status = str(filters.get('status', '') or '').strip().lower()
    if status and status != 'all':
        parts.append('Click status: %s' % status.replace('_', ' '))
    contact_status = str(filters.get('contact_status', '') or '').strip().upper()
    if contact_status in ('USED', 'UNUSED'):
        parts.append('Survey status: %s' % contact_status)
    for key, label in (('link_id', 'Link'), ('group_id', 'Group'),
                       ('job_id', 'Job ID'), ('search', 'Search')):
        value = str(filters.get(key, '') or '').strip()
        if value:
            parts.append('%s: %s' % (label, value))
    date_from = str(filters.get('date_from', '') or '').strip()
    date_to = str(filters.get('date_to', '') or '').strip()
    if date_from or date_to:
        parts.append('Date: %s – %s' % (date_from or '…', date_to or '…'))
    return '; '.join(parts)


def _reporting_period(filters):
    filters = filters or {}
    date_from = str(filters.get('date_from', '') or '').strip()
    date_to = str(filters.get('date_to', '') or '').strip()
    if date_from or date_to:
        return '%s – %s' % (date_from or 'beginning', date_to or 'present')
    return 'All dates'


# ---------------------------------------------------------------------------
# Summary dashboard
# ---------------------------------------------------------------------------
CARD_COLS = (2, 4, 6, 8)  # B-C, D-E, F-G, H-I


def _kpi_card(ws, row, col, label, value, sub):
    """One dashboard card: label / large number / context line."""
    ws.merge_cells(start_row=row, start_column=col,
                   end_row=row, end_column=col + 1)
    label_cell = ws.cell(row=row, column=col, value=label)
    label_cell.font = CARD_LABEL_FONT
    label_cell.alignment = CENTER
    ws.merge_cells(start_row=row + 1, start_column=col,
                   end_row=row + 1, end_column=col + 1)
    value_cell = ws.cell(row=row + 1, column=col, value=value)
    value_cell.font = CARD_VALUE_FONT
    value_cell.alignment = CENTER
    value_cell.number_format = COUNT_FORMAT
    ws.merge_cells(start_row=row + 2, start_column=col,
                   end_row=row + 2, end_column=col + 1)
    sub_cell = ws.cell(row=row + 2, column=col, value=sub)
    sub_cell.font = CARD_SUB_FONT
    sub_cell.alignment = CENTER
    for r in (row, row + 1, row + 2):
        for c in (col, col + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = CARD_FILL
            cell.border = THIN_BORDER
        edge = ws.cell(row=r, column=col).border
        ws.cell(row=r, column=col).border = Border(
            left=CARD_EDGE, right=edge.right, top=edge.top,
            bottom=edge.bottom)


def _kpi_card_rows(ws, row, card_rows):
    for cards in card_rows:
        for (label, value, sub), col in zip(cards, CARD_COLS):
            _kpi_card(ws, row, col, label, value, sub)
        ws.row_dimensions[row].height = 16
        ws.row_dimensions[row + 1].height = 30
        ws.row_dimensions[row + 2].height = 16
        row += 4  # one blank spacer row between card rows
    return row


def _rate_strip(ws, row, rates):
    """Banded strip of labelled percentage cells (real numeric values)."""
    for (label, _fraction), col in zip(rates, CARD_COLS):
        ws.merge_cells(start_row=row, start_column=col,
                       end_row=row, end_column=col + 1)
        label_cell = ws.cell(row=row, column=col, value=label)
        label_cell.font = STRIP_LABEL_FONT
        label_cell.alignment = CENTER
    for (_label, fraction), col in zip(rates, CARD_COLS):
        ws.merge_cells(start_row=row + 1, start_column=col,
                       end_row=row + 1, end_column=col + 1)
        value_cell = ws.cell(row=row + 1, column=col, value=fraction)
        value_cell.font = STRIP_VALUE_FONT
        value_cell.alignment = CENTER
        value_cell.number_format = PERCENT_FORMAT
    for r in (row, row + 1):
        for col in range(1, 11):
            ws.cell(row=r, column=col).fill = STRIP_FILL
    ws.row_dimensions[row].height = 16
    ws.row_dimensions[row + 1].height = 26
    return row + 2


def _comparison_table(ws, row, per_campaign, totals, reminders_sent,
                      link_lifecycle):
    """Campaign comparison as a real Excel Table + distinct total row."""
    headers = ['Campaign', 'Recipients', 'Sent', 'Delivered', 'Opened',
               'Clicked', 'Completed', 'Incomplete', 'Reminder Eligible',
               'Reminders Sent']
    header_row = row
    for idx, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=idx, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = HEADER_BORDER
    ws.row_dimensions[row].height = 28
    row += 1
    for item in per_campaign:
        agg = item['agg']
        values = [agg['recipients'], agg['sent'], agg['delivered'],
                  agg['opened'], agg['clicked'], agg['completed'],
                  agg['incomplete'], agg['eligible'], item['reminders_sent']]
        name_cell = ws.cell(row=row, column=1, value=item['campaign'].name)
        name_cell.font = BODY_FONT
        name_cell.alignment = LEFT_WRAP
        if link_lifecycle:
            _internal_link(name_cell, 'Recipient Lifecycle', target='A2')
            name_cell.font = LINK_FONT
        for idx, value in enumerate(values, start=2):
            cell = ws.cell(row=row, column=idx, value=value)
            cell.font = BODY_FONT
            cell.alignment = CENTER
            cell.number_format = COUNT_FORMAT
        ws.row_dimensions[row].height = 20
        row += 1
    last_data_row = row - 1
    _add_table(ws, 'CampaignComparison',
               'A%d:J%d' % (header_row, last_data_row))
    # Combined total: deliberately outside the table so sorting the
    # table never moves it.
    ws.cell(row=row, column=1, value='Combined Total').font = TOTAL_FONT
    total_values = [totals['recipients'], totals['sent'], totals['delivered'],
                    totals['opened'], totals['clicked'], totals['completed'],
                    totals['incomplete'], totals['eligible'], reminders_sent]
    for idx, value in enumerate(total_values, start=2):
        cell = ws.cell(row=row, column=idx, value=value)
        cell.font = TOTAL_FONT
        cell.alignment = CENTER
        cell.number_format = COUNT_FORMAT
    ws.cell(row=row, column=1).alignment = LEFT_WRAP
    for col in range(1, 11):
        ws.cell(row=row, column=col).fill = NAVY_FILL
    ws.row_dimensions[row].height = 22
    return row + 2, header_row, last_data_row  # blank spacer row after total


def _add_summary_charts(ws, anchor_row, header_row, last_data_row, totals):
    """Performance column chart + lifecycle doughnut. Returns next free row."""
    # Hidden helper cells for the doughnut (labels + live values).
    ws.cell(row=anchor_row, column=12, value='Completed')
    ws.cell(row=anchor_row, column=13, value=totals['completed'])
    ws.cell(row=anchor_row + 1, column=12, value='Incomplete')
    ws.cell(row=anchor_row + 1, column=13, value=totals['incomplete'])
    ws.column_dimensions['L'].hidden = True
    ws.column_dimensions['M'].hidden = True

    perf = BarChart()
    perf.type = 'col'
    perf.style = 10
    perf.title = 'Campaign Performance'
    perf.y_axis.title = 'Count'
    perf.height = 6.5
    perf.width = 11.5
    perf.gapWidth = 150
    perf.legend.position = 'b'
    for col, _series in ((3, 'Sent'), (4, 'Delivered'), (7, 'Completed')):
        perf.add_data(
            Reference(ws, min_col=col, min_row=header_row,
                      max_row=last_data_row),
            titles_from_data=True)
    perf.set_categories(
        Reference(ws, min_col=1, min_row=header_row + 1,
                  max_row=last_data_row))
    for series, color in zip(perf.series,
                             (CHART_SENT, CHART_DELIVERED, CHART_COMPLETED)):
        series.graphicalProperties.solidFill = color
    ws.add_chart(perf, 'B%d' % anchor_row)

    doughnut = DoughnutChart()
    doughnut.style = 10
    doughnut.title = 'Recipient Lifecycle'
    doughnut.height = 6.5
    doughnut.width = 9
    doughnut.holeSize = 50
    doughnut.legend = None
    doughnut.add_data(
        Reference(ws, min_col=13, min_row=anchor_row,
                  max_row=anchor_row + 1),
        titles_from_data=False)
    doughnut.set_categories(
        Reference(ws, min_col=12, min_row=anchor_row,
                  max_row=anchor_row + 1))
    series = doughnut.series[0]
    series.dLbls = DataLabelList()
    series.dLbls.showVal = True
    series.dLbls.showCatName = True
    for idx, color in enumerate((CHART_COMPLETED, CHART_INCOMPLETE)):
        point = DataPoint(idx=idx)
        point.graphicalProperties.solidFill = color
        series.dPt.append(point)
    ws.add_chart(doughnut, 'F%d' % anchor_row)
    return anchor_row + 14


def _report_content_nav(ws, row, ncols, links):
    """Two-column navigation buttons to the enabled data sheets."""
    positions = [(2, 4), (6, 8)]  # (start col, end col) pairs
    idx = 0
    while idx < len(links):
        for start_col, end_col in positions:
            if idx >= len(links):
                break
            sheet, label = links[idx]
            ws.merge_cells(start_row=row, start_column=start_col,
                           end_row=row, end_column=end_col)
            cell = ws.cell(row=row, column=start_col, value='Go to %s' % label)
            cell.font = LINK_FONT
            cell.alignment = CENTER
            cell.fill = NAV_FILL
            cell.border = THIN_BORDER
            _internal_link(cell, sheet, target='A2')
            idx += 1
        ws.row_dimensions[row].height = 24
        row += 1
        if idx < len(links):
            row += 1  # spacer between button rows
    return row + 1


def build_summary_sheet(wb, campaigns, per_campaign, filters, sections=None):
    from django.utils import timezone
    sections = sections or {}
    ws = wb.active
    ws.title = 'Summary'
    ws.sheet_properties.tabColor = '0F172A'
    ncols = 10
    widths = [3, 22, 14, 14, 14, 14, 14, 14, 14, 3]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    campaigns_label = ', '.join(c.name for c in campaigns)
    generated_label = timezone.localtime(timezone.now()).strftime('%d %b %Y, %I:%M %p %Z')
    row = _summary_title(ws, ncols, campaigns_label, generated_label,
                         _reporting_period(filters), _filters_summary(filters))

    combined_rows = [r for item in per_campaign for r in item['rows']]
    totals = _aggregate(
        combined_rows,
        sent=sum(item['agg']['sent'] for item in per_campaign),
        delivered=sum(item['agg']['delivered'] for item in per_campaign),
    )
    reminders_sent = sum(item['reminders_sent'] for item in per_campaign)

    row = _section_heading(ws, row, ncols, 'KEY PERFORMANCE INDICATORS')
    row = _kpi_card_rows(ws, row, [
        [('Total Recipients', totals['recipients'], 'Contacts in selection'),
         ('Emails Sent', totals['sent'], 'INITIAL + REMINDER'),
         ('Delivered', totals['delivered'], 'Excludes bounces'),
         ('Opened', totals['opened'], 'Unique recipients')],
        [('Clicked', totals['clicked'], 'Unique recipients'),
         ('Completed', totals['completed'], 'ODK submitted (USED)'),
         ('Incomplete / UNUSED', totals['incomplete'], 'Not yet submitted'),
         ('Reminder Eligible', totals['eligible'], 'Awaiting reminders')],
    ])

    row = _section_heading(ws, row, ncols, 'ENGAGEMENT RATES')
    delivery_rate = totals['delivered'] / totals['sent'] if totals['sent'] else 0.0
    row = _rate_strip(ws, row, [
        ('Delivery Rate', delivery_rate),
        ('Open Rate', totals['open_rate'] / 100),
        ('Click Rate', totals['click_rate'] / 100),
        ('Completion Rate', totals['completion_rate'] / 100),
    ])
    row += 1
    for (label, _value), col in zip(
            [('Reminders Sent', reminders_sent),
             ('Reminder Suppressed', totals['suppressed'])], (2, 4)):
        ws.merge_cells(start_row=row, start_column=col,
                       end_row=row, end_column=col + 1)
        label_cell = ws.cell(row=row, column=col, value=label)
        label_cell.font = STRIP_LABEL_FONT
        label_cell.alignment = CENTER
    for (_label, value), col in zip(
            [('Reminders Sent', reminders_sent),
             ('Reminder Suppressed', totals['suppressed'])], (2, 4)):
        ws.merge_cells(start_row=row + 1, start_column=col,
                       end_row=row + 1, end_column=col + 1)
        value_cell = ws.cell(row=row + 1, column=col, value=value)
        value_cell.font = STRIP_VALUE_FONT
        value_cell.alignment = CENTER
        value_cell.number_format = COUNT_FORMAT
    ws.row_dimensions[row].height = 16
    ws.row_dimensions[row + 1].height = 26
    row += 3

    row = _section_heading(ws, row, ncols, 'CAMPAIGN COMPARISON')
    row, cmp_header, cmp_last = _comparison_table(
        ws, row, per_campaign, totals, reminders_sent,
        link_lifecycle=_section_on(sections, 'lifecycle'))

    if totals['recipients']:
        row = _add_summary_charts(ws, row, cmp_header, cmp_last, totals)
        row += 1

    if _section_on(sections, 'campaign_info'):
        row = _section_heading(ws, row, ncols, 'CAMPAIGN INFORMATION')
        for campaign in campaigns:
            groups = ', '.join(g.name for g in campaign.groups.all()) or '—'
            status = (campaign.get_status_display()
                      if hasattr(campaign, 'get_status_display')
                      else campaign.status)
            name_cell = ws.cell(row=row, column=2, value=campaign.name)
            name_cell.font = BODY_FONT
            name_cell.alignment = LEFT_WRAP
            name_cell.border = THIN_BORDER
            if _section_on(sections, 'lifecycle'):
                _internal_link(name_cell, 'Recipient Lifecycle', target='A2')
                name_cell.font = LINK_FONT
            ws.merge_cells(start_row=row, start_column=3, end_row=row,
                           end_column=4)
            status_cell = ws.cell(row=row, column=3, value=status)
            status_cell.font = BODY_FONT
            status_cell.alignment = LEFT_WRAP
            status_cell.border = THIN_BORDER
            ws.merge_cells(start_row=row, start_column=5, end_row=row,
                           end_column=9)
            groups_cell = ws.cell(row=row, column=5, value=groups)
            groups_cell.font = BODY_FONT
            groups_cell.alignment = LEFT_WRAP
            for col in (3, 4, 5, 6, 7, 8, 9):
                ws.cell(row=row, column=col).border = THIN_BORDER
            ws.row_dimensions[row].height = 20
            row += 1
        row += 1

    nav_links = []
    if _section_on(sections, 'lifecycle'):
        nav_links.append(('Recipient Lifecycle', 'Recipient Lifecycle'))
    if _section_on(sections, 'clicks'):
        nav_links.append(('Click Activity', 'Click Activity'))
    if _section_on(sections, 'incomplete'):
        nav_links.append(('Incomplete - UNUSED', 'Incomplete — Follow-up'))
    if _section_on(sections, 'reminders'):
        nav_links.append(('Reminder Activity', 'Reminder Activity'))
    if nav_links:
        row = _section_heading(ws, row, ncols, 'REPORT CONTENT')
        row = _report_content_nav(ws, row, ncols, nav_links)

    note = ('Definitions match the Campaign Report: recipients are contacts assigned '
            'to (or messaged by) the campaign; sent/delivered count INITIAL/REMINDER '
            'messages (TEST messages excluded); opened/clicked count recipients; '
            'clicks are first-party link clicks; completed means ODK survey '
            'submitted (USED). Reminder eligibility never depends on opened/clicked state. '
            'All figures on this sheet are calculated from the same filtered selection '
            'shown on the data sheets.')
    ws.merge_cells(start_row=row, start_column=2, end_row=row + 2,
                   end_column=ncols - 1)
    note_cell = ws.cell(row=row, column=2, value=note)
    note_cell.font = SMALL_FONT
    note_cell.alignment = LEFT_WRAP

    _finish_sheet_common(ws, 'Summary')
    ws.print_area = 'A1:J%d' % (row + 2)
    ws.freeze_panes = 'A7'
    return ws


def build_info_sheet(wb, campaigns, filters):
    """Minimal identity sheet used when the Summary section is off."""
    from django.utils import timezone
    ws = wb.create_sheet('Report Info')
    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 90
    generated_label = timezone.localtime(timezone.now()).strftime('%d %b %Y, %I:%M %p %Z')
    ws.merge_cells('A1:B1')
    title = ws.cell(row=1, column=1, value='CAMPAIGN REPORT')
    title.font = TITLE_FONT
    title.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[1].height = 30
    for col in (1, 2):
        ws.cell(row=1, column=col).fill = NAVY_FILL
    meta = [
        ('Generated', generated_label),
        ('Campaigns', ', '.join(c.name for c in campaigns)),
        ('Reporting period', _reporting_period(filters)),
        ('Applied filters', _filters_summary(filters)),
    ]
    for idx, (label, value) in enumerate(meta, start=2):
        ws.cell(row=idx, column=1, value=label).font = META_LABEL_FONT
        ws.cell(row=idx, column=2, value=value).font = BODY_FONT
    _finish_sheet_common(ws, 'Report Info')
    return ws


# ---------------------------------------------------------------------------
# Interactive data sheets (real Excel Tables + conditional formatting)
# ---------------------------------------------------------------------------
def _data_sheet(wb, title, table_name, headers, widths, data_rows,
                date_cols=(), wrap_cols=(), center_cols=(),
                freeze='A3', back_target='Summary', cf_specs=()):
    """Styled data sheet: nav row, Excel Table, CF, freeze, print setup.

    cf_specs: ((column_1based, 'MATCH TEXT'), ...) applied to the data
    body with the shared CF_STYLES treatments.
    """
    ncols = len(headers)
    last_col = get_column_letter(ncols)
    ws = wb.create_sheet(title)
    back_label = ('← Back to Summary' if back_target == 'Summary'
                  else '← Back to Report Info')
    back = ws.cell(row=1, column=1, value=back_label)
    back.font = LINK_FONT
    back.alignment = LEFT_MID
    _internal_link(back, back_target, target='A1')
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=ncols)
    sheet_title = ws.cell(row=1, column=2, value=title.upper())
    sheet_title.font = SHEET_TITLE_FONT
    sheet_title.alignment = RIGHT_ALIGN
    ws.row_dimensions[1].height = 22

    _header_row(ws, 2, headers, widths)

    if not data_rows:
        ws.merge_cells(start_row=3, start_column=1, end_row=3,
                       end_column=ncols)
        cell = ws.cell(row=3, column=1, value=EMPTY_NOTE)
        cell.font = NOTE_FONT
        cell.alignment = CENTER
        last_row = 2
    else:
        for r_idx, values in enumerate(data_rows, start=3):
            tall = False
            for c_idx, value in enumerate(values, start=1):
                cell = ws.cell(row=r_idx, column=c_idx)
                if c_idx in date_cols:
                    if value is None:
                        cell.value = '—'
                    else:
                        cell.value = excel_dt(value)
                        cell.number_format = DATETIME_FORMAT
                    cell.font = BODY_FONT
                    cell.alignment = CENTER
                else:
                    cell.value = value
                    cell.font = BODY_FONT
                    if c_idx in wrap_cols:
                        cell.alignment = LEFT_WRAP
                        if isinstance(value, str) and len(value) > 60:
                            tall = True
                    elif c_idx in center_cols:
                        cell.alignment = CENTER
                    else:
                        cell.alignment = LEFT_MID
            ws.row_dimensions[r_idx].height = 30 if tall else 20
        last_row = 2 + len(data_rows)

    ref = 'A2:%s%d' % (last_col, max(last_row, 2))
    _add_table(ws, table_name, ref)
    cf_last = max(last_row, 3)
    for col_idx, text in cf_specs:
        col_letter = get_column_letter(col_idx)
        _cf_equal(ws, '%s3:%s%d' % (col_letter, col_letter, cf_last), text)

    ws.freeze_panes = freeze
    ws.print_title_rows = '2:2'
    _finish_sheet_common(ws, title)
    ws.print_area = 'A1:%s%d' % (last_col, max(last_row, 3))
    return ws


def build_lifecycle_sheet(wb, per_campaign, back_target='Summary'):
    headers = ['Campaign', 'Name', 'Phone', 'Email', 'Email Sent', 'Sent At',
               'Opened', 'Opened At', 'Clicked', 'Click Count',
               'First Clicked At', 'Tracking Token', 'Completed',
               'Completed At', 'ODK Submission ID', 'Reminder Eligible',
               'Reminder Status', 'Reminder Reason']
    widths = [22, 26, 16, 32, 13, 20, 11, 20, 11, 14, 21, 20, 13, 20, 22, 18, 17, 34]
    data = []
    for item in per_campaign:
        for r in item['rows']:
            data.append([
                item['campaign'].name,
                r['name'], r['phone'], r['email'],
                'YES' if r['has_sent'] else '—', r['email_sent_dt'],
                'YES' if r['has_opened'] else '—', r['email_opened_dt'],
                'YES' if r['has_clicked'] else '—', r['total_clicks'],
                r['first_click_dt'], r['tracking_token'] or '—',
                'COMPLETED' if r['contact_status'] == 'USED' else 'INCOMPLETE',
                r['completed_dt'], r['odk_submission_id'] or '—',
                'ELIGIBLE' if r['reminder_eligible'] else 'SUPPRESSED',
                reminder_status_label(r), r['reminder_reason'],
            ])
    return _data_sheet(
        wb, 'Recipient Lifecycle', 'RecipientLifecycle', headers, widths,
        data, date_cols=(6, 8, 11, 14), wrap_cols=(2, 4, 18),
        center_cols=(5, 6, 7, 8, 9, 10, 11, 13, 14, 16),
        freeze='C3', back_target=back_target,
        cf_specs=((5, 'YES'), (7, 'YES'), (9, 'YES'),
                  (13, 'COMPLETED'), (13, 'INCOMPLETE'),
                  (16, 'ELIGIBLE'), (16, 'SUPPRESSED')))


def build_clicks_sheet(wb, per_campaign, back_target='Summary'):
    headers = ['Campaign', 'Contact', 'Phone', 'Email', 'Tracking Token',
               'Clicked At', 'Device', 'Browser', 'Operating System',
               'Referrer', 'Click Type']
    widths = [22, 26, 16, 32, 20, 20, 14, 16, 20, 34, 15]
    data = []
    for item in per_campaign:
        for e in item['clicks']:
            data.append([
                item['campaign'].name, e['contact_name'], e['contact_phone'],
                e['contact_email'], e['tracking_token'], e['clicked_dt'],
                e['device_type'], e['browser'], e['operating_system'],
                e['referrer'] or '—', e['click_type'],
            ])
    return _data_sheet(
        wb, 'Click Activity', 'ClickActivity', headers, widths, data,
        date_cols=(6,), wrap_cols=(2, 4, 10), center_cols=(6, 7, 8, 11),
        back_target=back_target, cf_specs=((11, 'Suspected Bot'),))


def build_incomplete_sheet(wb, per_campaign, back_target='Summary'):
    headers = ['Campaign', 'Name', 'Phone', 'Email', 'Sent', 'Opened',
               'Clicked', 'Last Click', 'Completed', 'Reminder Eligible',
               'Reminder Status', 'Reminder Reason', 'Tracking Token']
    widths = [22, 26, 16, 32, 10, 11, 11, 20, 13, 18, 17, 34, 20]
    data = []
    for item in per_campaign:
        for r in item['rows']:
            if r['contact_status'] == 'USED':
                continue
            data.append([
                item['campaign'].name, r['name'], r['phone'], r['email'],
                'YES' if r['has_sent'] else '—',
                'YES' if r['has_opened'] else '—',
                'YES' if r['has_clicked'] else '—',
                r['last_click_dt'], 'INCOMPLETE',
                'ELIGIBLE' if r['reminder_eligible'] else 'SUPPRESSED',
                reminder_status_label(r), r['reminder_reason'],
                r['tracking_token'] or '—',
            ])
    return _data_sheet(
        wb, 'Incomplete - UNUSED', 'IncompleteUnused', headers, widths,
        data, date_cols=(8,), wrap_cols=(2, 4, 12),
        center_cols=(5, 6, 7, 8, 9, 10),
        back_target=back_target,
        cf_specs=((5, 'YES'), (6, 'YES'), (7, 'YES'),
                  (9, 'INCOMPLETE'),
                  (10, 'ELIGIBLE'), (10, 'SUPPRESSED')))


def build_reminders_sheet(wb, per_campaign, back_target='Summary'):
    headers = ['Campaign', 'Recipient', 'Email', 'Reminder Sequence',
               'Status', 'Sent At', 'Suppression / Skip Reason']
    widths = [22, 26, 32, 20, 14, 20, 44]
    data = []
    for item in per_campaign:
        for m in item['reminders']:
            data.append([
                item['campaign'].name, m['contact_name'], m['contact_email'],
                m['sequence'], m['status_label'], m['sent_dt'],
                m['reason'] or '—',
            ])
    return _data_sheet(
        wb, 'Reminder Activity', 'ReminderActivity', headers, widths, data,
        date_cols=(6,), wrap_cols=(2, 3, 7), center_cols=(4, 5, 6),
        back_target=back_target,
        cf_specs=((5, 'Sent'), (5, 'Failed')))


def build_campaign_workbook(campaigns, filters=None, sections=None,
                            scope_all=False):
    """Build the multi-campaign workbook.

    Returns (bytes, filename). Every sheet after Summary shows only rows
    from the same filtered selection the Summary aggregates are built
    from, so the numbers always reconcile.
    """
    from openpyxl import Workbook
    filters = filters or {}
    sections = sections or {}

    per_campaign = []
    for campaign in campaigns:
        rows = get_recipient_lifecycle_rows(campaign, filters)
        selection_ids = [r['contact_id'] for r in rows]
        # Message-based sent/delivered with the Campaign Report's exact
        # status sets, scoped to the same filtered selection (two
        # indexed COUNT queries per campaign).
        from apps.campaigns.models import CampaignMessage
        msg_qs = campaign.messages.filter(
            message_type__in=[
                CampaignMessage.MessageType.INITIAL,
                CampaignMessage.MessageType.REMINDER,
            ],
            contact_id__in=selection_ids,
        )
        sent_msgs = msg_qs.filter(status__in=[
            CampaignMessage.Status.SENT, CampaignMessage.Status.DELIVERED,
            CampaignMessage.Status.OPENED, CampaignMessage.Status.CLICKED,
        ]).count()
        delivered_msgs = msg_qs.filter(status__in=[
            CampaignMessage.Status.DELIVERED,
            CampaignMessage.Status.OPENED, CampaignMessage.Status.CLICKED,
        ]).count()
        clicks, _total = get_click_activity_rows(
            campaign,
            link_id=(filters.get('link_id') or None),
            date_from=(filters.get('date_from') or None),
            date_to=(filters.get('date_to') or None),
            contact_ids=selection_ids,
            limit=None,
        )
        reminders = get_reminder_activity_rows(
            campaign, contact_ids=selection_ids)
        reminders_sent = sum(1 for m in reminders if m['sent_dt'])
        per_campaign.append({
            'campaign': campaign,
            'rows': rows,
            'clicks': clicks,
            'reminders': reminders,
            'reminders_sent': reminders_sent,
            'agg': _aggregate(rows, sent=sent_msgs, delivered=delivered_msgs),
        })

    wb = Workbook()
    wb.properties.title = 'Campaign Report — %s' % (
        ', '.join(c.name for c in campaigns))
    wb.properties.creator = 'IRIS Marketing'
    wb.properties.lastModifiedBy = 'IRIS Marketing'
    wb.properties.description = (
        'Campaign performance workbook. Filters: %s'
        % _filters_summary(filters))

    want_summary = _section_on(sections, 'summary')
    if want_summary:
        build_summary_sheet(wb, campaigns, per_campaign, filters, sections)
    else:
        # Drop the default sheet; keep a minimal identity sheet so the
        # workbook always states what it contains.
        wb.remove(wb.active)
        build_info_sheet(wb, campaigns, filters)

    back_target = 'Summary' if want_summary else 'Report Info'
    if _section_on(sections, 'lifecycle'):
        build_lifecycle_sheet(wb, per_campaign, back_target=back_target)
    if _section_on(sections, 'clicks'):
        build_clicks_sheet(wb, per_campaign, back_target=back_target)
    if _section_on(sections, 'incomplete'):
        build_incomplete_sheet(wb, per_campaign, back_target=back_target)
    if _section_on(sections, 'reminders'):
        build_reminders_sheet(wb, per_campaign, back_target=back_target)

    wb.active = 0
    buf = BytesIO()
    wb.save(buf)
    filename = build_report_filename(
        [c.name for c in campaigns],
        lifecycle=str(filters.get('lifecycle', 'all') or 'all'),
        search=str(filters.get('search', '') or ''),
        scope_all=scope_all,
    )
    return buf.getvalue(), filename
