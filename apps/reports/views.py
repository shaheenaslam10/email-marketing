from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status, permissions
from apps.campaigns.models import Campaign, CampaignMessage
from apps.contacts.models import Contact
from apps.groups.models import ContactGroup
from apps.reminders.models import ReminderConfiguration
from apps.tracking.models import EmailEvent
from django.db.models import Q
from .services import (
    get_campaign_full_report, get_recipient_lifecycle_rows,
    get_click_activity_rows,
)
from .exporters import export_campaign_xlsx, export_campaign_csv, export_campaign_pdf


class CampaignReportView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request, pk):
        try:
            campaign = Campaign.objects.get(pk=pk)
        except Campaign.DoesNotExist:
            return Response({'error': 'Campaign not found'}, status=status.HTTP_404_NOT_FOUND)

        report_data = get_campaign_full_report(campaign)
        return Response(report_data)


class CampaignReportMessagesView(APIView):
    """
    Drill-down API returning contact messages for a campaign filtered by delivery status.
    Supports status in: Sent, Delivered, Soft Bounce, Hard Bounce, Failed, Opened, Clicked.
    """
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request, pk):
        try:
            campaign = Campaign.objects.get(pk=pk)
        except Campaign.DoesNotExist:
            return Response({'error': 'Campaign not found'}, status=status.HTTP_404_NOT_FOUND)

        status_filter = request.GET.get('status', 'Sent').strip()
        search_query = request.GET.get('search', '').strip()
        stage_filter = request.GET.get('stage', '').strip()

        messages = campaign.messages.select_related('contact').all()

        norm_status = status_filter.lower().replace(' ', '_')
        if norm_status == 'sent':
            messages = messages.filter(status__in=[
                CampaignMessage.Status.SENT, CampaignMessage.Status.DELIVERED,
                CampaignMessage.Status.OPENED, CampaignMessage.Status.CLICKED
            ])
        elif norm_status == 'delivered':
            messages = messages.filter(status__in=[
                CampaignMessage.Status.DELIVERED, CampaignMessage.Status.OPENED,
                CampaignMessage.Status.CLICKED
            ])
        elif norm_status == 'soft_bounce':
            messages = messages.filter(status=CampaignMessage.Status.SOFT_BOUNCE)
        elif norm_status == 'hard_bounce':
            messages = messages.filter(status=CampaignMessage.Status.HARD_BOUNCE)
        elif norm_status == 'failed':
            messages = messages.filter(status=CampaignMessage.Status.FAILED)
        elif norm_status == 'opened':
            messages = messages.filter(opened_at__isnull=False)
        elif norm_status == 'clicked':
            messages = messages.filter(clicked_at__isnull=False)
        elif norm_status != 'all':
            messages = messages.filter(status__iexact=status_filter)

        if stage_filter:
            if stage_filter.lower() == 'initial':
                messages = messages.filter(message_type=CampaignMessage.MessageType.INITIAL)
            elif 'reminder' in stage_filter.lower():
                parts = stage_filter.split()
                if len(parts) > 1 and parts[1].isdigit():
                    messages = messages.filter(message_type=CampaignMessage.MessageType.REMINDER, reminder_sequence=int(parts[1]))

        if search_query:
            messages = messages.filter(
                Q(contact__first_name__icontains=search_query) |
                Q(contact__last_name__icontains=search_query) |
                Q(contact__email__icontains=search_query) |
                Q(contact__job_id__icontains=search_query) |
                Q(to_email__icontains=search_query)
            )

        messages = messages.order_by('-sent_at', '-id')
        total_count = messages.count()

        results = []
        for m in messages[:300]:
            c = m.contact
            stage_name = 'Initial Email' if m.message_type == CampaignMessage.MessageType.INITIAL else f'Reminder {m.reminder_sequence}'
            results.append({
                'id': m.id,
                'contact_id': c.id if c else None,
                'name': c.name if c else (m.to_email or 'Unknown Contact'),
                'email': m.to_email or (c.email if c else ''),
                'job_id': c.job_id if c else '',
                'stage': stage_name,
                'status': m.status,
                'contact_status': c.status if c else 'UNKNOWN',
                'sent_at': m.sent_at.isoformat() if m.sent_at else None,
                'delivered_at': m.delivered_at.isoformat() if m.delivered_at else None,
                'opened_at': m.opened_at.isoformat() if m.opened_at else None,
                'error_message': m.error_message or m.skip_reason or '',
            })

        return Response({
            'status': status_filter,
            'total_count': total_count,
            'messages': results
        })


class CampaignReportLinkRecipientsView(APIView):
    """
    Recipient-Level Link Tracking Report (Requirements 6, 8, 9, 14).
    Returns detailed recipient activity with multi-criteria filters, search,
    custom pagination (200 | 400 | 600 | 1000 | All), and CSV / Excel export.
    """
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request, pk):
        try:
            campaign = Campaign.objects.get(pk=pk)
        except Campaign.DoesNotExist:
            return Response({'error': 'Campaign not found'}, status=status.HTTP_404_NOT_FOUND)

        import csv
        import io
        from django.http import HttpResponse

        # Shared row builder: the JSON view, CSV/XLSX exports and the
        # Excel workbook export all consume these same rows.
        rows = get_recipient_lifecycle_rows(campaign, {
            'status': request.GET.get('status', 'all'),
            'contact_status': request.GET.get('contact_status', 'all'),
            'link_id': request.GET.get('link_id', ''),
            'group_id': request.GET.get('group_id', ''),
            'job_id': request.GET.get('job_id', ''),
            'search': request.GET.get('search', ''),
            'date_from': request.GET.get('date_from', ''),
            'date_to': request.GET.get('date_to', ''),
            'lifecycle': request.GET.get('lifecycle', 'all'),
        })
        total_count = len(rows)
        export_fmt = request.GET.get('export', '').strip().lower()

        # Handle CSV export
        if export_fmt == 'csv':
            response = HttpResponse(content_type='text/csv')
            response['Content-Disposition'] = f'attachment; filename="campaign_{campaign.id}_link_tracking.csv"'
            writer = csv.writer(response)
            writer.writerow([
                'Name', 'Email', 'Phone', 'JobID', 'Survey Status',
                'Email Sent', 'Email Opened', 'Link Status', 'Tracking Token',
                'First Click', 'Last Click', 'Total Clicks', 'Human Clicks',
                'Bot Clicks', 'Click Type', 'Completed At', 'ODK Submission',
                'Reminder Eligible', 'Reminder Reason',
                'Link Name', 'Short URL', 'Original Destination'
            ])
            for r in rows:
                writer.writerow([
                    r['name'], r['email'], r['phone'], r['job_id'], r['contact_status'],
                    r['email_sent_at'], r['email_opened_at'], r['link_status'], r['tracking_token'],
                    r['first_click'], r['last_click'], r['total_clicks'], r['human_clicks'],
                    r['bot_clicks'], r['click_type'], r['completed_at'], r['odk_submission_id'],
                    'Yes' if r['reminder_eligible'] else 'No', r['reminder_reason'],
                    r['link_name'], r['short_url'], r['original_url']
                ])
            return response

        # Handle Excel (XLSX) export
        if export_fmt in ('xlsx', 'excel'):
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Link Tracking Report"

            headers = [
                'Name', 'Email', 'Phone', 'JobID', 'Survey Status',
                'Email Sent', 'Email Opened', 'Link Status', 'Tracking Token',
                'First Click', 'Last Click', 'Total Clicks', 'Human Clicks',
                'Bot Clicks', 'Click Type', 'Completed At', 'ODK Submission',
                'Reminder Eligible', 'Reminder Reason',
                'Link Name', 'Short URL', 'Original Destination'
            ]
            ws.append(headers)

            # Style header row
            header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Append rows
            for r in rows:
                ws.append([
                    r['name'], r['email'], r['phone'], r['job_id'], r['contact_status'],
                    r['email_sent_at'], r['email_opened_at'], r['link_status'], r['tracking_token'],
                    r['first_click'], r['last_click'], r['total_clicks'], r['human_clicks'],
                    r['bot_clicks'], r['click_type'], r['completed_at'], r['odk_submission_id'],
                    'Yes' if r['reminder_eligible'] else 'No', r['reminder_reason'],
                    r['link_name'], r['short_url'], r['original_url']
                ])

            # Auto adjust column widths
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            response = HttpResponse(
                output.read(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = f'attachment; filename="campaign_{campaign.id}_link_tracking.xlsx"'
            return response

        # Pagination: 200 | 400 | 600 | 1000 | All
        page_size_str = request.GET.get('page_size', '200').strip().lower()
        if page_size_str == 'all':
            page_size = max(total_count, 1)
        else:
            try:
                page_size = int(page_size_str)
                if page_size not in (200, 400, 600, 1000):
                    page_size = 200
            except ValueError:
                page_size = 200

        try:
            page = max(1, int(request.GET.get('page', 1)))
        except ValueError:
            page = 1

        total_pages = (total_count + page_size - 1) // page_size if total_count > 0 else 1
        start_idx = (page - 1) * page_size
        end_idx = start_idx + page_size
        paginated_rows = rows[start_idx:end_idx]

        return Response({
            'campaign_id': campaign.id,
            'campaign_name': campaign.name,
            'total_count': total_count,
            'page': page,
            'page_size': page_size if page_size_str != 'all' else 'all',
            'total_pages': total_pages,
            'recipients': paginated_rows
        })


class CampaignReportLinkClicksView(APIView):
    """
    Recipient Click Activity: per-click events for recipient email links.
    Unions legacy root-level LinkClickEvent rows with RECIPIENT-type /c/
    visits, newest first. Shareable (anonymous) visits are excluded here;
    they belong to the shareable-links report. IP addresses are stored but
    excluded from this API (privacy; staff can use Django admin).
    """
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request, pk):
        try:
            campaign = Campaign.objects.get(pk=pk)
        except Campaign.DoesNotExist:
            return Response({'error': 'Campaign not found'}, status=status.HTTP_404_NOT_FOUND)

        from apps.tracking.models import ShortenedLink

        try:
            limit = int(request.GET.get('limit', 50))
        except (TypeError, ValueError):
            limit = 50
        limit = max(1, min(limit, 200))

        link_id = (request.GET.get('link_id') or '').strip()

        # Shared row builder (legacy + RECIPIENT /c/ union, newest first).
        events, total_events = get_click_activity_rows(
            campaign, link_id=link_id or None, limit=limit)

        links = [
            {'id': sl.id, 'name': sl.link_name or sl.original_url[:40]}
            for sl in ShortenedLink.objects.filter(
                campaign=campaign).order_by('link_name', 'id')
        ]

        return Response({
            'campaign_id': campaign.id,
            'campaign_name': campaign.name,
            'total_events': total_events,
            'links': links,
            'events': events,
        })


class CampaignReportShareableLinksView(APIView):
    """
    Shareable / Anonymous Link Activity Report.
    Covers SHAREABLE-type CampaignTrackingLink rows (/c/<token>/) and
    their CampaignLinkClickEvent visits. Fully separate from
    recipient-level email tracking: anonymous visits carry no contact
    attribution, so no "unique clicks" metric is offered (no visitor
    identity is tracked). RECIPIENT-type /c/ links are excluded here;
    they appear in the recipient Link Clicks report instead.
    Visitor IP addresses are stored but intentionally excluded from this
    API (privacy; available to staff via Django admin).
    """
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request, pk):
        try:
            campaign = Campaign.objects.get(pk=pk)
        except Campaign.DoesNotExist:
            return Response({'error': 'Campaign not found'}, status=status.HTTP_404_NOT_FOUND)

        from apps.tracking.models import CampaignTrackingLink, CampaignLinkClickEvent
        from apps.tracking.utils import build_campaign_short_url

        links = list(campaign.tracking_links.filter(
            link_type=CampaignTrackingLink.LinkType.SHAREABLE
        ).order_by('-created_at'))

        link_id = (request.GET.get('link_id') or '').strip()
        selected_link = None
        if link_id.isdigit():
            selected_link = next((l for l in links if l.id == int(link_id)), None)

        try:
            limit = int(request.GET.get('limit', 50))
        except (TypeError, ValueError):
            limit = 50
        limit = max(1, min(limit, 200))

        events_qs = CampaignLinkClickEvent.objects.filter(
            campaign=campaign,
            link__link_type=CampaignTrackingLink.LinkType.SHAREABLE,
        )
        if selected_link:
            events_qs = events_qs.filter(link=selected_link)
        events = list(events_qs.select_related('link').order_by('-clicked_at', '-id')[:limit])

        links_payload = []
        for l in links:
            links_payload.append({
                'id': l.id,
                'name': l.name or 'Untitled link',
                # Derived from the current base (see link-clicks).
                'short_url': build_campaign_short_url(l.tracking_token),
                'destination_url': l.resolve_destination(),
                'uses_campaign_default': not (l.destination_url or '').strip(),
                'is_active': l.is_active,
                'click_count': l.click_count,
                'human_click_count': l.human_click_count,
                'bot_click_count': l.bot_click_count,
                'first_clicked_at': l.first_clicked_at.isoformat() if l.first_clicked_at else None,
                'last_clicked_at': l.last_clicked_at.isoformat() if l.last_clicked_at else None,
            })

        def click_type_label(ct):
            return {
                'HUMAN': 'Human',
                'SUSPECTED_BOT': 'Suspected Bot',
                'UNKNOWN': 'Unknown',
            }.get(ct, 'Unknown')

        events_payload = [{
            'clicked_at': e.clicked_at.isoformat() if e.clicked_at else None,
            'link_id': e.link_id,
            'link_name': ((e.link.name if e.link else '') or 'Untitled link'),
            'click_type': click_type_label(e.click_type),
            'browser': e.browser or 'Other',
            'operating_system': e.operating_system or 'Other',
            'device_type': e.device_type or 'Unknown',
            'referrer': e.referrer or '',
        } for e in events]

        return Response({
            'campaign_id': campaign.id,
            'campaign_name': campaign.name,
            'campaign_destination_url': campaign.destination_url or '',
            'totals': {
                'links': len(links),
                'active_links': sum(1 for l in links if l.is_active),
                'total_clicks': sum(l.click_count for l in links),
                'human_clicks': sum(l.human_click_count for l in links),
                'bot_clicks': sum(l.bot_click_count for l in links),
            },
            'links': links_payload,
            'recent_events': events_payload,
            'recent_limit': limit,
        })


class CampaignExportView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request, pk, fmt):
        try:
            campaign = Campaign.objects.get(pk=pk)
        except Campaign.DoesNotExist:
            return Response({'error': 'Campaign not found'}, status=status.HTTP_404_NOT_FOUND)

        fmt_lower = fmt.lower()
        if fmt_lower == 'xlsx':
            return export_campaign_xlsx(campaign)
        elif fmt_lower == 'csv':
            return export_campaign_csv(campaign, {
                'lifecycle': request.GET.get('lifecycle', 'all'),
                'status': request.GET.get('status', 'all'),
                'contact_status': request.GET.get('contact_status', 'all'),
                'link_id': request.GET.get('link_id', ''),
                'group_id': request.GET.get('group_id', ''),
                'job_id': request.GET.get('job_id', ''),
                'date_from': request.GET.get('date_from', ''),
                'date_to': request.GET.get('date_to', ''),
            })
        elif fmt_lower == 'pdf':
            return export_campaign_pdf(campaign)
        return Response({'error': f'Unsupported export format {fmt}'}, status=status.HTTP_400_BAD_REQUEST)


class ReportXlsxExportView(APIView):
    """
    Multi-campaign Excel workbook export (the "Export Report" dialog).
    Accepts JSON: campaign_ids [...], filters {...}, sections {...}.
    Same IsAuthenticated access as every other campaign report endpoint.
    """
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        from django.http import HttpResponse
        from .services import get_visible_campaigns
        from .workbook import build_campaign_workbook

        data = request.data if isinstance(request.data, dict) else {}
        scope = str(data.get('scope', '') or '').strip().lower()
        raw_ids = data.get('campaign_ids', data.get('campaigns', []))
        scope_all = scope == 'all' or (
            isinstance(raw_ids, str) and raw_ids.strip().lower() == 'all')

        # Same visibility as the /campaigns/ listing page: submitted ids
        # outside the visible set are rejected, never exported.
        visible = get_visible_campaigns(request.user)
        if scope_all:
            campaigns = list(visible.prefetch_related('groups'))
            if not campaigns:
                return Response({'error': 'No campaigns available to export.'},
                                status=status.HTTP_400_BAD_REQUEST)
        else:
            if isinstance(raw_ids, int):
                raw_ids = [raw_ids]
            try:
                ids = [int(v) for v in (raw_ids or [])]
            except (TypeError, ValueError):
                return Response({'error': 'campaign_ids must be a list of campaign ids.'},
                                status=status.HTTP_400_BAD_REQUEST)
            # Dedupe, preserve order
            ids = list(dict.fromkeys(ids))
            if not ids:
                return Response({'error': 'Select at least one campaign.'},
                                status=status.HTTP_400_BAD_REQUEST)

            # No campaign-count cap: per-campaign cost is ~10 indexed
            # queries plus its own rows, so "All Campaigns" stays linear.
            campaigns = list(visible.filter(pk__in=ids).prefetch_related('groups'))
            found = {c.id for c in campaigns}
            missing = [v for v in ids if v not in found]
            if missing:
                existing = set(Campaign.objects.filter(
                    pk__in=missing).values_list('pk', flat=True))
                forbidden = [v for v in missing if v in existing]
                if forbidden:
                    return Response(
                        {'error': f'Not authorized for campaign(s): {forbidden}'},
                        status=status.HTTP_403_FORBIDDEN)
                return Response({'error': f'Unknown campaign(s): {missing}'},
                                status=status.HTTP_400_BAD_REQUEST)

        filters = data.get('filters') or {}
        if not isinstance(filters, dict):
            return Response({'error': 'filters must be an object.'},
                            status=status.HTTP_400_BAD_REQUEST)
        sections = data.get('sections') or {}
        if not isinstance(sections, dict):
            return Response({'error': 'sections must be an object.'},
                            status=status.HTTP_400_BAD_REQUEST)

        content, filename = build_campaign_workbook(
            campaigns, filters=filters, sections=sections, scope_all=scope_all)
        response = HttpResponse(
            content,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


class DashboardStatsView(APIView):
    """Global system dashboard metrics (Section 77 & 78)."""
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        total_contacts = Contact.objects.count()
        used_contacts = Contact.objects.filter(status=Contact.UsageStatus.USED).count()
        unused_contacts = Contact.objects.filter(status=Contact.UsageStatus.UNUSED).count()
        total_groups = ContactGroup.objects.count()
        total_campaigns = Campaign.objects.count()
        active_reminders = ReminderConfiguration.objects.filter(enabled=True, campaign__status=Campaign.Status.ACTIVE).count()

        emails_sent = CampaignMessage.objects.filter(sent_at__isnull=False).count()
        emails_delivered = CampaignMessage.objects.filter(delivered_at__isnull=False).count()
        emails_opened = CampaignMessage.objects.filter(opened_at__isnull=False).count()
        emails_clicked = CampaignMessage.objects.filter(clicked_at__isnull=False).count()
        unsubscribed = Contact.objects.filter(unsubscribed=True).count()
        bounced = Contact.objects.filter(email_status__in=[Contact.EmailStatus.HARD_BOUNCE, Contact.EmailStatus.SOFT_BOUNCE]).count()

        # Active campaigns summary table (Section 78)
        active_campaigns = []
        for c in Campaign.objects.filter(status__in=[Campaign.Status.ACTIVE, Campaign.Status.SCHEDULED])[:10]:
            groups = c.groups.all()
            contacts = Contact.objects.filter(groups__in=groups).distinct()
            recipients = contacts.count()
            c_used = contacts.filter(status=Contact.UsageStatus.USED).count()
            c_unused = contacts.filter(status=Contact.UsageStatus.UNUSED).count()
            comp_rate = round((c_used / recipients * 100), 1) if recipients > 0 else 0.0

            next_rem = None
            if hasattr(c, 'reminder_config') and c.reminder_config.enabled:
                next_rem = c.reminder_config.next_run_at

            active_campaigns.append({
                'id': c.id,
                'name': c.name,
                'recipients': recipients,
                'used': c_used,
                'unused': c_unused,
                'completion_rate': comp_rate,
                'next_reminder': next_rem,
                'status': c.status
            })

        return Response({
            'total_contacts': total_contacts,
            'used_contacts': used_contacts,
            'unused_contacts': unused_contacts,
            'completion_rate': round(used_contacts / total_contacts * 100, 1) if total_contacts > 0 else 0.0,
            'total_groups': total_groups,
            'total_campaigns': total_campaigns,
            'active_reminder_campaigns': active_reminders,
            'emails_sent': emails_sent,
            'delivered': emails_delivered,
            'opened': emails_opened,
            'clicked': emails_clicked,
            'unsubscribed': unsubscribed,
            'bounced': bounced,
            'active_campaigns': active_campaigns
        })
