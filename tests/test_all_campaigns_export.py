"""All / multiple campaigns export from the /campaigns/ listing page.

Same POST /api/campaigns/report/xlsx/ endpoint and same shared workbook
services as the individual-campaign Export Report; this module pins the
multi-campaign workflows: scope 'all', selected campaigns, date ranges,
sections, combined totals, visibility rules and filenames. Asserts on
real XLSX bytes via openpyxl throughout.
"""
import json
from datetime import timedelta
from io import BytesIO

from django.test import TestCase, Client
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.campaigns.models import Campaign, CampaignMessage
from apps.campaigns.views import CampaignViewSet
from apps.contacts.models import Contact
from apps.groups.models import ContactGroup
from apps.odk.models import ODKConnection, ODKProject, ODKForm, ODKFieldMapping
from apps.odk.services import run_odk_sync
from apps.reminders.models import ReminderConfiguration
from apps.reminders.services import execute_reminder_cycle, launch_initial_campaign
from apps.reports.services import get_campaign_full_report, get_visible_campaigns
from apps.sandbox.models import MockODKSubmission
from apps.senders.models import Sender
from apps.tracking.models import CampaignTrackingLink, TrackingToken

ODK_URL = 'https://odk.example.com/f/FAKE123?st=FakeToken'
HUMAN_UA = ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 '
            '(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36')
URL = '/api/campaigns/report/xlsx/'
EXPECTED_COLUMNS = ['Campaign', 'Recipients', 'Sent', 'Delivered', 'Opened',
                    'Clicked', 'Completed', 'Incomplete', 'Reminder Eligible',
                    'Reminders Sent']


class AllCampaignsExportTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username='all_admin', email='all@example.com', password='password123')
        self.client = Client()
        self.client.force_login(self.user)
        self.anon = Client()

        self.sender = Sender.objects.create(
            name='Survey Team', email='survey@marketing.iriscommunications.cloud',
            provider_type=Sender.ProviderType.SANDBOX, is_active=True)
        self.group1 = ContactGroup.objects.create(name='All V1 Group')
        self.group2 = ContactGroup.objects.create(name='All V2 Group')

        odk_conn = ODKConnection.objects.create(
            name='Mock ODK', base_url='http://localhost:8000/sandbox/api/odk',
            is_mock_sandbox=True, status=ODKConnection.Status.CONNECTED)
        odk_project = ODKProject.objects.create(
            connection=odk_conn, odk_id=1, name='Pulse 2026')
        self.odk_form = ODKForm.objects.create(
            project=odk_project, odk_xml_form_id='pulse_v1', name='Pulse_V1')
        ODKFieldMapping.objects.create(form=self.odk_form, job_id_field='job_id')

        self.v1 = self._make_campaign('Pulse V1', self.group1)
        self.v2 = self._make_campaign('Pulse V2', self.group2)
        ReminderConfiguration.objects.create(
            campaign=self.v1, enabled=True, interval_value=2,
            interval_unit=ReminderConfiguration.Unit.DAYS,
            max_reminders=5, sync_odk_before_send=True, stop_when_used=True)

        self.a = self._make_contact('Amy Done', 'amy@example.com', 'VXA', self.group1)
        self.b = self._make_contact('Ben Clicked', 'ben@example.com', 'VXB', self.group1)
        self.c = self._make_contact('Cat Quiet', 'cat@example.com', 'VXC', self.group1)
        self.d = self._make_contact('Dee Opened', 'dee@example.com', 'VXD', self.group2)
        self.e = self._make_contact('Eli Silent', 'eli@example.com', 'VXE', self.group2)

        self.assertEqual(launch_initial_campaign(self.v1), 3)
        self.assertEqual(launch_initial_campaign(self.v2), 2)

        self._open(self.v1, self.a)
        self._open(self.v1, self.b)
        self._open(self.v2, self.d)
        self._click(self.v1, self.a)
        self._click(self.v1, self.b)
        # B's clicks happened 11 days ago (fixed past for date-range tests).
        past = timezone.now() - timedelta(days=11)
        CampaignTrackingLink.objects.filter(
            campaign=self.v1, contact=self.b,
            link_type=CampaignTrackingLink.LinkType.RECIPIENT).update(
                first_clicked_at=past, last_clicked_at=past)

        MockODKSubmission.objects.create(
            submission_id='uuid:sub-amy', project_id='1', form_id='pulse_v1',
            job_id='VXA', respondent_email='amy@example.com')
        run_odk_sync(self.odk_form)
        self.a.refresh_from_db()
        self.assertEqual(self.a.status, Contact.UsageStatus.USED)

        execute_reminder_cycle(self.v1, manual_trigger=True)

    def _make_campaign(self, name, group):
        campaign = Campaign.objects.create(
            name=name, campaign_type=Campaign.Type.SURVEY_REMINDER,
            status=Campaign.Status.ACTIVE, sender=self.sender,
            odk_form=self.odk_form, subject='Survey for {{job_id}}',
            html_content='<p>Dear {{first_name}},</p><p>{{survey_tracking_url}}</p>',
            destination_url=ODK_URL)
        campaign.groups.add(group)
        return campaign

    def _make_contact(self, name, email, job_id, group):
        c = Contact.objects.create(
            name=name, first_name=name.split(' ')[0], email=email,
            phone_number='0300000000', job_id=job_id,
            status=Contact.UsageStatus.UNUSED)
        group.contacts.add(c)
        return c

    def _initial_msg(self, campaign, contact):
        return CampaignMessage.objects.get(
            campaign=campaign, contact=contact,
            message_type=CampaignMessage.MessageType.INITIAL)

    def _open(self, campaign, contact):
        msg = self._initial_msg(campaign, contact)
        token = TrackingToken.objects.get(message=msg).token
        resp = self.anon.get('/t/open/%s/' % token)
        self.assertEqual(resp.status_code, 200)

    def _click(self, campaign, contact):
        link = CampaignTrackingLink.objects.get(
            campaign=campaign, contact=contact,
            link_type=CampaignTrackingLink.LinkType.RECIPIENT)
        resp = self.anon.get('/c/%s/' % link.tracking_token, HTTP_USER_AGENT=HUMAN_UA)
        self.assertEqual(resp.status_code, 302)

    # -- helpers ------------------------------------------------------
    def _export(self, payload, client=None):
        client = client or self.client
        return client.post(URL, data=json.dumps(payload), content_type='application/json')

    def _wb(self, response):
        self.assertEqual(response.status_code, 200)
        return load_workbook(BytesIO(response.content))

    def _table_rows(self, ws):
        header_row = None
        for row in ws.iter_rows(values_only=False):
            if row[0].value == 'Campaign':
                header_row = row[0].row
                break
        self.assertIsNotNone(header_row, 'table headers missing')
        headers = [ws.cell(row=header_row, column=i).value
                   for i in range(1, ws.max_column + 1)]
        while headers and headers[-1] is None:
            headers.pop()
        rows = []
        for values in ws.iter_rows(min_row=header_row + 1,
                                   max_col=len(headers), values_only=True):
            if values[0] is None:
                continue
            rows.append(dict(zip(headers, values)))
        return headers, rows

    def _kpis(self, ws):
        # Dashboard cards/strips: the value sits directly below the label.
        kpis = {}
        for row in ws.iter_rows():
            for cell in row:
                if (isinstance(cell.value, str) and cell.value in (
                        'Total Recipients', 'Completed', 'Incomplete / UNUSED')
                        and cell.value not in kpis):
                    kpis[cell.value] = ws.cell(
                        row=cell.row + 1, column=cell.column).value
        return kpis

    def _meta(self, ws):
        meta = {}
        for row in ws.iter_rows(min_row=1, max_row=8, values_only=False):
            cells = list(row)
            for i, cell in enumerate(cells):
                if isinstance(cell.value, str) and cell.value.endswith(':'):
                    for nxt in cells[i + 1:]:
                        if nxt.value is not None:
                            meta[cell.value] = nxt.value
                            break
        return meta

    def _assert_empty_note(self, ws):
        notes = [row[0] for row in
                 ws.iter_rows(min_col=1, max_col=1, values_only=True)]
        self.assertIn('No records match the selected filters.', notes)

    def _campaign_table(self, ws):
        header_row = None
        for row in ws.iter_rows(values_only=False):
            if row[0].value == 'Campaign' and row[1].value == 'Recipients':
                header_row = row[0].row
                break
        self.assertIsNotNone(header_row)
        headers = [ws.cell(row=header_row, column=i).value for i in range(1, 11)]
        rows, combined = [], None
        r = header_row + 1
        while True:
            name = ws.cell(row=r, column=1).value
            if name is None:
                break
            values = [ws.cell(row=r, column=i).value for i in range(1, 11)]
            if name == 'Combined Total':
                combined = dict(zip(headers, values))
                break
            rows.append(dict(zip(headers, values)))
            r += 1
        return headers, rows, combined

    # -- listing page ---------------------------------------------------
    def test_01_listing_page_has_export_reports(self):
        resp = self.client.get('/campaigns/')
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Export Reports')
        self.assertContains(resp, 'id="reports-modal"')
        self.assertContains(resp, 'onclick="openReportsDialog()"')
        self.assertContains(resp, 'onclick="generateReportsExport()"')
        self.assertContains(resp, 'id="reports-lifecycle"')
        self.assertContains(resp, '/api/campaigns/report/xlsx/')
        # Login wall intact for anonymous visitors.
        self.assertEqual(self.anon.get('/campaigns/').status_code, 302)

    def test_01b_reports_dialog_wiring(self):
        # Same regression guard as the individual dialog: modal markup
        # must render inside a block and every JS-referenced control
        # must exist.
        resp = self.client.get('/campaigns/')
        html = resp.content.decode('utf-8')
        for element_id in (
                'reports-modal', 'reports-error', 'reports-campaign-picker',
                'reports-campaign-search', 'reports-campaign-list',
                'reports-campaign-count', 'reports-lifecycle',
                'reports-job-id', 'reports-date-from',
                'reports-date-to', 'reports-sec-summary',
                'reports-sec-campaign-info', 'reports-sec-lifecycle',
                'reports-sec-clicks', 'reports-sec-incomplete',
                'reports-sec-reminders', 'reports-generate',
                'reports-spinner', 'reports-icon', 'reports-generate-label'):
            self.assertIn('id="%s"' % element_id, html, element_id)
        self.assertEqual(html.count('type="radio" name="reports-scope"'), 2)
        self.assertNotIn('reports-search', html)
        import re
        for used in set(re.findall(r"getElementById\('([^']+)'\)", html)):
            if used.startswith('reports-'):
                self.assertIn('id="%s"' % used, html, used)

    # -- scope: all -------------------------------------------------------
    def test_02_scope_all_exports_everything(self):
        resp = self._export({'scope': 'all'})
        wb = self._wb(resp)
        self.assertIn('Campaign_Report_All_Campaigns_', resp['Content-Disposition'])
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual(len(rows), 5)
        self.assertEqual({r['Campaign'] for r in rows}, {'Pulse V1', 'Pulse V2'})
        kpis = self._kpis(wb['Summary'])
        self.assertEqual(kpis['Total Recipients'], 5)
        self.assertEqual(kpis['Completed'], 1)
        self.assertEqual(kpis['Incomplete / UNUSED'], 4)

    def test_03_scope_all_string_ids_variant(self):
        resp = self._export({'campaign_ids': 'all'})
        wb = self._wb(resp)
        self.assertIn('Campaign_Report_All_Campaigns_', resp['Content-Disposition'])
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual(len(rows), 5)

    def test_04_scope_all_wins_over_ids(self):
        resp = self._export({'scope': 'all', 'campaign_ids': [self.v1.id]})
        wb = self._wb(resp)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual({r['Campaign'] for r in rows}, {'Pulse V1', 'Pulse V2'})

    # -- scope: selected ----------------------------------------------------
    def test_05_single_selected_campaign(self):
        resp = self._export({'campaign_ids': [self.v2.id]})
        wb = self._wb(resp)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual({r['Email'] for r in rows},
                         {'dee@example.com', 'eli@example.com'})
        self._assert_empty_note(wb['Click Activity'])
        self.assertIn('Campaign_Report_Pulse_V2_', resp['Content-Disposition'])

    def test_06_multiple_selected_campaigns(self):
        resp = self._export({'campaign_ids': [self.v1.id, self.v2.id]})
        wb = self._wb(resp)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual(len(rows), 5)
        _h, clicks = self._table_rows(wb['Click Activity'])
        self.assertEqual({e['Email'] for e in clicks},
                         {'amy@example.com', 'ben@example.com'})
        self.assertIn('Campaign_Report_Pulse_V1_Pulse_V2_',
                      resp['Content-Disposition'])

    # -- date range (existing click-date semantics) ------------------------------
    def test_07_date_range(self):
        today = timezone.localdate()
        date_from = (today - timedelta(days=6)).isoformat()
        resp = self._export({'scope': 'all', 'filters': {'date_from': date_from}})
        wb = self._wb(resp)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        # Only Amy clicked recently; Ben's clicks are 11 days old and
        # everyone else never clicked (no first click => out of range).
        self.assertEqual([r['Email'] for r in rows], ['amy@example.com'])
        _h, clicks = self._table_rows(wb['Click Activity'])
        self.assertEqual([e['Email'] for e in clicks], ['amy@example.com'])
        self.assertIn(date_from, self._meta(wb['Summary'])['Reporting period:'])

        date_to = (today - timedelta(days=10)).isoformat()
        resp = self._export({'scope': 'all', 'filters': {'date_to': date_to}})
        wb = self._wb(resp)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual([r['Email'] for r in rows], ['ben@example.com'])

    # -- filters ---------------------------------------------------------------
    def test_08_lifecycle_and_search_filters(self):
        resp = self._export({'scope': 'all',
                             'filters': {'lifecycle': 'incomplete'}})
        wb = self._wb(resp)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual({r['Email'] for r in rows},
                         {'ben@example.com', 'cat@example.com',
                          'dee@example.com', 'eli@example.com'})
        resp = self._export({'scope': 'all', 'filters': {'search': 'dee'}})
        wb = self._wb(resp)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual([r['Email'] for r in rows], ['dee@example.com'])
        resp = self._export({'scope': 'all', 'filters': {'job_id': 'VXE'}})
        wb = self._wb(resp)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual([r['Email'] for r in rows], ['eli@example.com'])

    # -- sections -----------------------------------------------------------------
    def test_09_section_selection(self):
        resp = self._export(
            {'scope': 'all',
             'sections': {'summary': True, 'campaign_info': False,
                          'lifecycle': False, 'clicks': False,
                          'incomplete': True, 'reminders': False}})
        wb = self._wb(resp)
        self.assertEqual(wb.sheetnames, ['Summary', 'Incomplete - UNUSED'])
        titles = [wb['Summary'].cell(row=r, column=1).value
                  for r in range(1, wb['Summary'].max_row + 1)]
        self.assertNotIn('CAMPAIGN INFORMATION', titles)
        # Default keeps the Campaign Information block on Summary.
        resp = self._export({'campaign_ids': [self.v1.id]})
        wb = self._wb(resp)
        titles = [wb['Summary'].cell(row=r, column=1).value
                  for r in range(1, wb['Summary'].max_row + 1)]
        self.assertIn('CAMPAIGN INFORMATION', titles)

    # -- combined totals -------------------------------------------------------------
    def test_10_combined_totals(self):
        api1 = get_campaign_full_report(self.v1)
        api2 = get_campaign_full_report(self.v2)
        resp = self._export({'scope': 'all'})
        wb = self._wb(resp)
        headers, table, combined = self._campaign_table(wb['Summary'])
        self.assertEqual(headers, EXPECTED_COLUMNS)
        self.assertEqual(len(table), 2)
        self.assertEqual(
            combined['Recipients'],
            api1['kpi']['initial_eligible'] + api2['kpi']['initial_eligible'])
        self.assertEqual(
            combined['Sent'], api1['kpi']['sent_count'] + api2['kpi']['sent_count'])
        self.assertEqual(
            combined['Completed'],
            api1['survey_conversion']['used_count'] + api2['survey_conversion']['used_count'])
        self.assertEqual(combined['Incomplete'], 4)
        v2row = [r for r in table if r['Campaign'] == 'Pulse V2'][0]
        self.assertEqual(v2row['Recipients'], 2)
        self.assertEqual(v2row['Clicked'], 0)

    # -- permissions --------------------------------------------------------------------
    def test_11_permissions(self):
        resp = self._export({'scope': 'all'}, client=self.anon)
        self.assertIn(resp.status_code, (401, 403))
        resp = self._export({'campaign_ids': [self.v1.id]}, client=self.anon)
        self.assertIn(resp.status_code, (401, 403))
        resp = self._export({'campaign_ids': [424242]})
        self.assertEqual(resp.status_code, 400)
        # Visibility helper mirrors the listing exactly.
        visible = get_visible_campaigns(self.user)
        self.assertEqual(set(visible.values_list('pk', flat=True)),
                         set(CampaignViewSet.queryset.values_list('pk', flat=True)))
        self.assertEqual(get_visible_campaigns(None).count(), 0)
        from django.contrib.auth.models import AnonymousUser
        self.assertEqual(get_visible_campaigns(AnonymousUser()).count(), 0)

    # -- empty ------------------------------------------------------------------------------
    def test_12_empty_results(self):
        Campaign.objects.all().delete()
        resp = self._export({'scope': 'all'})
        self.assertEqual(resp.status_code, 400)
        self.assertIn('No campaigns', resp.json()['error'])

    def test_13_empty_filtered_selection(self):
        resp = self._export({'scope': 'all',
                             'filters': {'search': 'nobody-matches-this'}})
        wb = self._wb(resp)
        kpis = self._kpis(wb['Summary'])
        self.assertEqual(kpis['Total Recipients'], 0)
        _h, headers_rows, combined = self._campaign_table(wb['Summary'])
        self.assertEqual(combined['Recipients'], 0)
        for sheet in ('Recipient Lifecycle', 'Click Activity',
                      'Incomplete - UNUSED', 'Reminder Activity'):
            self._assert_empty_note(wb[sheet])
