"""Excel workbook export: multi-campaign professional report.

Seeds through REAL paths (launch -> /t/open/ -> /c/ -> ODK sandbox sync
-> reminder cycle) and asserts on real XLSX bytes: sheets, headers,
values, summary-vs-API parity, filter consistency, existing-export
compatibility, permissions and filenames.

Person C (never opened + never clicked + incomplete) must stay
reminder-eligible and appear on the Incomplete - UNUSED sheet.
"""
import json
from datetime import datetime
from io import BytesIO

from django.test import TestCase, Client
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.campaigns.models import Campaign, CampaignMessage
from apps.contacts.models import Contact
from apps.groups.models import ContactGroup
from apps.odk.models import ODKConnection, ODKProject, ODKForm, ODKFieldMapping
from apps.odk.services import run_odk_sync
from apps.reminders.models import ReminderConfiguration
from apps.reminders.services import execute_reminder_cycle, launch_initial_campaign
from apps.reports.services import build_report_filename, get_campaign_full_report
from apps.sandbox.models import MockODKSubmission
from apps.senders.models import Sender
from apps.tracking.models import CampaignTrackingLink, TrackingToken

ODK_URL = 'https://odk.example.com/f/FAKE123?st=FakeToken'
HUMAN_UA = ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 '
            '(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36')
URL = '/api/campaigns/report/xlsx/'

EXPECTED_SHEETS = ['Summary', 'Recipient Lifecycle', 'Click Activity',
                   'Incomplete - UNUSED', 'Reminder Activity']


class ExcelReportExportTests(TestCase):
    # -- seeding ------------------------------------------------------
    def setUp(self):
        self.user = User.objects.create_user(
            username='xlsx_admin', email='xlsx@example.com', password='password123')
        self.client = Client()
        self.client.force_login(self.user)
        self.anon = Client()

        self.sender = Sender.objects.create(
            name='Survey Team', email='survey@marketing.iriscommunications.cloud',
            provider_type=Sender.ProviderType.SANDBOX, is_active=True)
        self.group1 = ContactGroup.objects.create(name='Pulse V1 Group')
        self.group2 = ContactGroup.objects.create(name='Pulse V2 Group')

        odk_conn = ODKConnection.objects.create(
            name='Mock ODK', base_url='http://localhost:8000/sandbox/api/odk',
            is_mock_sandbox=True, status=ODKConnection.Status.CONNECTED)
        odk_project = ODKProject.objects.create(
            connection=odk_conn, odk_id=1, name='Pulse 2026')
        self.odk_form = ODKForm.objects.create(
            project=odk_project, odk_xml_form_id='pulse_v1', name='Pulse_V1')
        ODKFieldMapping.objects.create(form=self.odk_form, job_id_field='job_id')

        self.v1 = self._make_campaign('Pulse V1', self.group1)
        self.v2 = self._make_campaign('Pulse V2', self.group2)
        ReminderConfiguration.objects.create(
            campaign=self.v1, enabled=True, interval_value=2,
            interval_unit=ReminderConfiguration.Unit.DAYS,
            max_reminders=5, sync_odk_before_send=True, stop_when_used=True)

        # V1 roster: A completed, B engaged+incomplete, C unengaged+incomplete,
        # D hard-bounced, E unsubscribed.
        self.a = self._make_contact('Alice Completed', 'alice@example.com', '03001111111', 'VX-A', self.group1)
        self.b = self._make_contact('Bob Engaged', 'bob@example.com', '03002222222', 'VX-B', self.group1)
        self.c = self._make_contact('Carol Quiet', 'carol@example.com', '03003333333', 'VX-C', self.group1)
        self.d = self._make_contact('Dan Bounced', 'dan@example.com', '03004444444', 'VX-D', self.group1)
        self.e = self._make_contact('Eve Optout', 'eve@example.com', '03005555555', 'VX-E', self.group1)
        # V2 roster: F opened+incomplete, G completed w/o open/click, H sent only.
        self.f = self._make_contact('Farah Opened', 'farah@example.com', '03006666666', 'VX-F', self.group2)
        self.g = self._make_contact('Gary Done', 'gary@example.com', '03007777777', 'VX-G', self.group2)
        self.h = self._make_contact('Hina Silent', 'hina@example.com', '03008888888', 'VX-H', self.group2)

        self.assertEqual(launch_initial_campaign(self.v1), 5)
        self.assertEqual(launch_initial_campaign(self.v2), 3)

        # D hard-bounces at delivery; E opts out after delivery.
        d_msg = self._initial_msg(self.v1, self.d)
        d_msg.status = CampaignMessage.Status.HARD_BOUNCE
        d_msg.sent_at = None
        d_msg.delivered_at = None
        d_msg.save(update_fields=['status', 'sent_at', 'delivered_at'])
        self.d.email_status = Contact.EmailStatus.HARD_BOUNCE
        self.d.save(update_fields=['email_status'])
        self.e.unsubscribed = True
        self.e.save(update_fields=['unsubscribed'])

        self._open(self.v1, self.a)
        self._open(self.v1, self.b)
        self._open(self.v2, self.f)
        self._click(self.v1, self.a)
        self._click(self.v1, self.b)

        MockODKSubmission.objects.create(
            submission_id='uuid:sub-A', project_id='1', form_id='pulse_v1',
            job_id='VX-A', respondent_email='alice@example.com')
        MockODKSubmission.objects.create(
            submission_id='uuid:sub-G', project_id='1', form_id='pulse_v1',
            job_id='VX-G', respondent_email='gary@example.com')
        run_odk_sync(self.odk_form)
        self.a.refresh_from_db()
        self.g.refresh_from_db()
        self.assertEqual(self.a.status, Contact.UsageStatus.USED)
        self.assertEqual(self.g.status, Contact.UsageStatus.USED)

        cycle = execute_reminder_cycle(self.v1, manual_trigger=True)
        self.assertEqual(cycle.eligible_count, 2)  # B and C only

    def _make_campaign(self, name, group):
        campaign = Campaign.objects.create(
            name=name, campaign_type=Campaign.Type.SURVEY_REMINDER,
            status=Campaign.Status.ACTIVE, sender=self.sender,
            odk_form=self.odk_form, subject='Survey for {{job_id}}',
            html_content='<p>Dear {{first_name}},</p><p>{{survey_tracking_url}}</p>',
            destination_url=ODK_URL)
        campaign.groups.add(group)
        return campaign

    def _make_contact(self, name, email, phone, job_id, group):
        c = Contact.objects.create(
            name=name, first_name=name.split(' ')[0], email=email,
            phone_number=phone, job_id=job_id,
            status=Contact.UsageStatus.UNUSED)
        group.contacts.add(c)
        return c

    def _initial_msg(self, campaign, contact):
        return CampaignMessage.objects.get(
            campaign=campaign, contact=contact,
            message_type=CampaignMessage.MessageType.INITIAL)

    def _open(self, campaign, contact):
        msg = self._initial_msg(campaign, contact)
        token = TrackingToken.objects.get(message=msg).token
        resp = self.anon.get('/t/open/%s/' % token)
        self.assertEqual(resp.status_code, 200)

    def _click(self, campaign, contact):
        link = CampaignTrackingLink.objects.get(
            campaign=campaign, contact=contact,
            link_type=CampaignTrackingLink.LinkType.RECIPIENT)
        resp = self.anon.get('/c/%s/' % link.tracking_token, HTTP_USER_AGENT=HUMAN_UA)
        self.assertEqual(resp.status_code, 302)
        return link

    # -- helpers ------------------------------------------------------
    def _export(self, campaign_ids, filters=None, sections=None, client=None):
        client = client or self.client
        payload = {'campaign_ids': campaign_ids}
        if filters is not None:
            payload['filters'] = filters
        if sections is not None:
            payload['sections'] = sections
        return client.post(URL, data=json.dumps(payload), content_type='application/json')

    def _wb(self, response):
        self.assertEqual(response.status_code, 200)
        return load_workbook(BytesIO(response.content))

    KPI_LABELS = ('Total Recipients', 'Emails Sent', 'Delivered', 'Opened',
                  'Clicked', 'Completed', 'Incomplete / UNUSED',
                  'Reminder Eligible', 'Reminders Sent', 'Reminder Suppressed')
    RATE_LABELS = ('Delivery Rate', 'Open Rate', 'Click Rate',
                   'Completion Rate')

    def _label_values(self, ws, labels):
        # Dashboard cards/strips: the value sits directly below the label.
        # First match wins (chart helper cells reuse plain words).
        found = {}
        for row in ws.iter_rows():
            for cell in row:
                if (isinstance(cell.value, str) and cell.value in labels
                        and cell.value not in found):
                    found[cell.value] = ws.cell(
                        row=cell.row + 1, column=cell.column).value
        return found

    def _kpis(self, ws):
        return self._label_values(ws, self.KPI_LABELS)

    def _rates(self, ws):
        return self._label_values(ws, self.RATE_LABELS)

    def _table_rows(self, ws):
        header_row = None
        for row in ws.iter_rows(values_only=False):
            if row[0].value == 'Campaign':
                header_row = row[0].row
                break
        self.assertIsNotNone(header_row, 'table headers missing')
        headers = [ws.cell(row=header_row, column=i).value
                   for i in range(1, ws.max_column + 1)]
        while headers and headers[-1] is None:
            headers.pop()
        rows = []
        for values in ws.iter_rows(min_row=header_row + 1,
                                   max_col=len(headers), values_only=True):
            if values[0] is None:
                continue
            rows.append(dict(zip(headers, values)))
        return headers, rows

    def _assert_empty_note(self, ws):
        notes = [row[0] for row in
                 ws.iter_rows(min_col=1, max_col=1, values_only=True)]
        self.assertIn('No records match the selected filters.', notes)

    def _chart_titles(self, ws):
        titles = []
        for chart in ws._charts:
            text = chart.title.tx.rich.p[0].r[0].t
            titles.append((type(chart).__name__, text))
        return titles

    def _meta(self, ws):
        meta = {}
        for row in ws.iter_rows(min_row=1, max_row=8, values_only=False):
            cells = list(row)
            for i, cell in enumerate(cells):
                if isinstance(cell.value, str) and cell.value.endswith(':'):
                    for nxt in cells[i + 1:]:
                        if nxt.value is not None:
                            meta[cell.value] = nxt.value
                            break
        return meta

    def _below(self, ws, label):
        for row in ws.iter_rows():
            for cell in row:
                if cell.value == label:
                    return ws.cell(row=cell.row + 1, column=cell.column)
        self.fail('label %r missing' % label)

    def _xlsx_xml(self, response, prefix):
        import zipfile
        zf = zipfile.ZipFile(BytesIO(response.content))
        return {name: zf.read(name).decode('utf-8')
                for name in zf.namelist() if name.startswith(prefix)}

    def _campaign_table(self, ws):
        header_row = None
        for row in ws.iter_rows(values_only=False):
            if row[0].value == 'Campaign' and row[1].value == 'Recipients':
                header_row = row[0].row
                break
        self.assertIsNotNone(header_row, 'campaign comparison table missing')
        headers = [ws.cell(row=header_row, column=i).value for i in range(1, 11)]
        rows, combined = [], None
        r = header_row + 1
        while True:
            name = ws.cell(row=r, column=1).value
            if name is None:
                break
            values = [ws.cell(row=r, column=i).value for i in range(1, 11)]
            if name == 'Combined Total':
                combined = dict(zip(headers, values))
                break
            rows.append(dict(zip(headers, values)))
            r += 1
        return rows, combined

    # 1. Single campaign ------------------------------------------------
    def test_01_single_campaign_export(self):
        resp = self._export([self.v1.id])
        self.assertEqual(
            resp['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        wb = self._wb(resp)
        self.assertEqual(wb.sheetnames, EXPECTED_SHEETS)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual(len(rows), 5)
        self.assertEqual({r['Campaign'] for r in rows}, {'Pulse V1'})
        kpis = self._kpis(wb['Summary'])
        self.assertEqual(kpis['Total Recipients'], 5)
        self.assertEqual(kpis['Completed'], 1)
        self.assertEqual(kpis['Incomplete / UNUSED'], 4)

    # 2 + 12. Multiple campaigns, one workbook --------------------------
    def test_02_multi_campaign_single_workbook(self):
        resp = self._export([self.v1.id, self.v2.id])
        wb = self._wb(resp)
        self.assertEqual(wb.sheetnames, EXPECTED_SHEETS)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual(len(rows), 8)
        self.assertEqual({r['Campaign'] for r in rows}, {'Pulse V1', 'Pulse V2'})
        table, combined = self._campaign_table(wb['Summary'])
        self.assertEqual(len(table), 2)
        self.assertEqual(combined['Recipients'], 8)
        self.assertEqual(combined['Completed'], 2)
        _h, inc = self._table_rows(wb['Incomplete - UNUSED'])
        self.assertEqual({r['Campaign'] for r in inc}, {'Pulse V1', 'Pulse V2'})
        _h, rem = self._table_rows(wb['Reminder Activity'])
        self.assertTrue(all(r['Campaign'] == 'Pulse V1' for r in rem))

    # 3. Current (page-level) filters -----------------------------------
    def test_03_status_and_contact_status_filters(self):
        resp = self._export(
            [self.v1.id],
            filters={'status': 'clicked', 'contact_status': 'UNUSED'})
        wb = self._wb(resp)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual([r['Email'] for r in rows], ['bob@example.com'])
        kpis = self._kpis(wb['Summary'])
        self.assertEqual(kpis['Total Recipients'], 1)

    # 4. Campaign filter --------------------------------------------------
    def test_04_campaign_subset(self):
        resp = self._export([self.v2.id])
        wb = self._wb(resp)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual(len(rows), 3)
        self.assertEqual({r['Campaign'] for r in rows}, {'Pulse V2'})
        table, combined = self._campaign_table(wb['Summary'])
        self.assertEqual(len(table), 1)
        self.assertEqual(table[0]['Campaign'], 'Pulse V2')

    # 5. Recipient search -------------------------------------------------
    def test_05_search_filter(self):
        resp = self._export([self.v1.id, self.v2.id], filters={'search': 'BOB'})
        wb = self._wb(resp)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual([r['Email'] for r in rows], ['bob@example.com'])
        resp = self._export([self.v1.id], filters={'search': 'VX-C'})
        wb = self._wb(resp)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual([r['Email'] for r in rows], ['carol@example.com'])

    # 6. Completed --------------------------------------------------------
    def test_06_completed_filter(self):
        resp = self._export(
            [self.v1.id, self.v2.id], filters={'lifecycle': 'completed'})
        wb = self._wb(resp)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual({r['Email'] for r in rows},
                         {'alice@example.com', 'gary@example.com'})
        self.assertTrue(all(r['Completed'] == 'COMPLETED' for r in rows))
        self._assert_empty_note(wb['Incomplete - UNUSED'])

    # 7. Incomplete / UNUSED (Person C included) ---------------------------
    def test_07_incomplete_filter_includes_person_c(self):
        resp = self._export([self.v1.id], filters={'lifecycle': 'incomplete'})
        wb = self._wb(resp)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual({r['Email'] for r in rows},
                         {'bob@example.com', 'carol@example.com',
                          'dan@example.com', 'eve@example.com'})
        _h, inc = self._table_rows(wb['Incomplete - UNUSED'])
        emails = {r['Email'] for r in inc}
        self.assertIn('carol@example.com', emails)  # Person C on follow-up sheet
        self.assertNotIn('alice@example.com', emails)
        self.assertIn('Incomplete', resp['Content-Disposition'])

    # 8. Clicked but not completed -----------------------------------------
    def test_08_clicked_not_completed(self):
        resp = self._export(
            [self.v1.id], filters={'lifecycle': 'clicked_not_completed'})
        wb = self._wb(resp)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual([r['Email'] for r in rows], ['bob@example.com'])

    # 9. Never opened + never clicked + incomplete --------------------------
    def test_09_unengaged_incomplete(self):
        resp = self._export(
            [self.v1.id], filters={'lifecycle': 'unengaged_incomplete'})
        wb = self._wb(resp)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual({r['Email'] for r in rows},
                         {'carol@example.com', 'dan@example.com', 'eve@example.com'})

    # 10. Reminder eligible -------------------------------------------------
    def test_10_reminder_eligible(self):
        resp = self._export(
            [self.v1.id, self.v2.id], filters={'lifecycle': 'reminder_eligible'})
        wb = self._wb(resp)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual({r['Email'] for r in rows},
                         {'bob@example.com', 'carol@example.com',
                          'farah@example.com', 'hina@example.com'})
        self.assertTrue(all(r['Reminder Eligible'] == 'ELIGIBLE' for r in rows))

    # 11. Reminder suppressed -----------------------------------------------
    def test_11_reminder_suppressed_with_reasons(self):
        resp = self._export(
            [self.v1.id], filters={'lifecycle': 'reminder_suppressed'})
        wb = self._wb(resp)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        reasons = {r['Email']: r['Reminder Reason'] for r in rows}
        self.assertEqual(reasons, {
            'alice@example.com': 'Survey completed',
            'dan@example.com': 'Email hard bounce',
            'eve@example.com': 'Unsubscribed',
        })
        self.assertTrue(all(r['Reminder Eligible'] == 'SUPPRESSED' for r in rows))

    # 12. Section toggles ----------------------------------------------------
    def test_12_section_toggles(self):
        resp = self._export(
            [self.v1.id],
            sections={'summary': False, 'lifecycle': True, 'clicks': False,
                      'incomplete': False, 'reminders': False})
        wb = self._wb(resp)
        self.assertEqual(wb.sheetnames, ['Report Info', 'Recipient Lifecycle'])
        resp = self._export(
            [self.v1.id],
            sections={'summary': True, 'lifecycle': False, 'clicks': False,
                      'incomplete': True, 'reminders': False})
        wb = self._wb(resp)
        self.assertEqual(wb.sheetnames, ['Summary', 'Incomplete - UNUSED'])

    # 13. Empty result ---------------------------------------------------------
    def test_13_empty_result(self):
        resp = self._export([self.v1.id], filters={'search': 'zzz-no-such-person'})
        wb = self._wb(resp)
        self.assertEqual(wb.sheetnames, EXPECTED_SHEETS)
        kpis = self._kpis(wb['Summary'])
        for label in ('Total Recipients', 'Emails Sent', 'Delivered', 'Opened',
                      'Clicked', 'Completed', 'Incomplete / UNUSED',
                      'Reminder Eligible', 'Reminders Sent', 'Reminder Suppressed'):
            self.assertEqual(kpis[label], 0, label)
        for sheet in ('Recipient Lifecycle', 'Click Activity',
                      'Incomplete - UNUSED', 'Reminder Activity'):
            self._assert_empty_note(wb[sheet])

    # 14. Long names / emails -----------------------------------------------------
    def test_14_long_names_and_emails(self):
        long_name = 'Alexandria ' + 'Montgomery-Fitzwilliam ' * 5 + 'Esquire'
        long_email = 'alexandria.montgomery.fitzwilliam.the.third.esquire.001@example-corporation.co.uk'
        c = self._make_contact(long_name, long_email, '03009999999', 'VX-LONG', self.group1)
        CampaignMessage.objects.create(
            campaign=self.v1, contact=c, to_email=long_email, subject='s',
            message_type=CampaignMessage.MessageType.INITIAL,
            status=CampaignMessage.Status.DELIVERED,
            sent_at=timezone.now(), delivered_at=timezone.now())
        resp = self._export([self.v1.id])
        wb = self._wb(resp)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        match = [r for r in rows if r['Email'] == long_email]
        self.assertEqual(len(match), 1)
        self.assertEqual(match[0]['Name'], long_name)

    # 15. Large dataset --------------------------------------------------------------
    def test_15_large_dataset(self):
        group = ContactGroup.objects.create(name='Scale Group')
        campaign = self._make_campaign('Pulse V3 Scale', group)
        now = timezone.now()
        contacts = [Contact(
            name='Scale User %04d' % i, first_name='Scale',
            email='scale%04d@example.com' % i, phone_number='0300000000',
            job_id='SCALE-%04d' % i, status=Contact.UsageStatus.UNUSED)
            for i in range(800)]
        Contact.objects.bulk_create(contacts)
        ids = list(Contact.objects.filter(job_id__startswith='SCALE-').values_list('id', flat=True))
        group.contacts.add(*ids)
        CampaignMessage.objects.bulk_create([
            CampaignMessage(
                campaign=campaign, contact_id=cid,
                to_email='scale%04d@example.com' % i, subject='s',
                message_type=CampaignMessage.MessageType.INITIAL,
                status=CampaignMessage.Status.DELIVERED,
                sent_at=now, delivered_at=now)
            for i, cid in enumerate(ids)])
        resp = self._export([campaign.id])
        wb = self._wb(resp)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual(len(rows), 800)
        kpis = self._kpis(wb['Summary'])
        self.assertEqual(kpis['Total Recipients'], 800)
        self.assertEqual(kpis['Incomplete / UNUSED'], 800)
        ws = wb['Recipient Lifecycle']
        self.assertEqual(ws.tables['RecipientLifecycle'].ref, 'A2:R802')
        ranges = [str(cf.sqref) for cf in ws.conditional_formatting]
        self.assertIn('M3:M802', ranges)  # bounded CF, no whole-column rules
        self.assertIn('P3:P802', ranges)

    # 16. Sheets -----------------------------------------------------------------------
    def test_16_default_sheets(self):
        resp = self._export([self.v1.id])
        wb = self._wb(resp)
        self.assertEqual(wb.sheetnames, EXPECTED_SHEETS)
        self.assertEqual(wb['Summary']['A1'].value, 'CAMPAIGN REPORT')
        self.assertIn('Campaign Report', wb.properties.title)

    # 17. Headers + workbook design ------------------------------------------------------
    def test_17_headers_and_design(self):
        resp = self._export([self.v1.id])
        wb = self._wb(resp)
        headers, _rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual(headers, ['Campaign', 'Name', 'Phone', 'Email',
                                   'Email Sent', 'Sent At', 'Opened', 'Opened At',
                                   'Clicked', 'Click Count', 'First Clicked At',
                                   'Tracking Token', 'Completed', 'Completed At',
                                   'ODK Submission ID', 'Reminder Eligible',
                                   'Reminder Status', 'Reminder Reason'])
        headers, _rows = self._table_rows(wb['Click Activity'])
        self.assertEqual(headers, ['Campaign', 'Contact', 'Phone', 'Email',
                                   'Tracking Token', 'Clicked At', 'Device',
                                   'Browser', 'Operating System', 'Referrer',
                                   'Click Type'])
        headers, _rows = self._table_rows(wb['Incomplete - UNUSED'])
        self.assertEqual(headers, ['Campaign', 'Name', 'Phone', 'Email', 'Sent',
                                   'Opened', 'Clicked', 'Last Click', 'Completed',
                                   'Reminder Eligible', 'Reminder Status',
                                   'Reminder Reason', 'Tracking Token'])
        headers, _rows = self._table_rows(wb['Reminder Activity'])
        self.assertEqual(headers, ['Campaign', 'Recipient', 'Email',
                                   'Reminder Sequence', 'Status', 'Sent At',
                                   'Suppression / Skip Reason'])
        # Interactive tables, freeze panes, print setup on every data sheet.
        expected = {
            'Recipient Lifecycle': ('RecipientLifecycle', 'C3'),
            'Click Activity': ('ClickActivity', 'A3'),
            'Incomplete - UNUSED': ('IncompleteUnused', 'A3'),
            'Reminder Activity': ('ReminderActivity', 'A3'),
        }
        for sheet, (table_name, freeze) in expected.items():
            ws = wb[sheet]
            self.assertEqual(ws.freeze_panes, freeze, sheet)
            self.assertIn(table_name, ws.tables, sheet)
            table = ws.tables[table_name]
            self.assertEqual(table.tableStyleInfo.name,
                             'TableStyleMedium2', sheet)
            self.assertTrue(table.tableStyleInfo.showRowStripes, sheet)
            self.assertIsNotNone(table.autoFilter, sheet)
            self.assertEqual(table.autoFilter.ref, table.ref, sheet)
            self.assertEqual(ws.print_title_rows, '$2:$2', sheet)
            self.assertEqual(ws.page_setup.orientation, 'landscape', sheet)
            self.assertEqual(ws['A1'].value, '\u2190 Back to Summary', sheet)
            self.assertEqual(ws['A1'].hyperlink.location,
                             "'Summary'!A1", sheet)

    # 18. Expected values -----------------------------------------------------------------
    def test_18_expected_values(self):
        resp = self._export([self.v1.id])
        wb = self._wb(resp)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        by_email = {r['Email']: r for r in rows}
        link = CampaignTrackingLink.objects.get(
            campaign=self.v1, contact=self.b,
            link_type=CampaignTrackingLink.LinkType.RECIPIENT)
        bob = by_email['bob@example.com']
        self.assertEqual(bob['Email Sent'], 'YES')
        self.assertIsInstance(bob['Sent At'], datetime)
        self.assertEqual(bob['Opened'], 'YES')
        self.assertEqual(bob['Clicked'], 'YES')
        self.assertEqual(bob['Click Count'], 1)
        self.assertEqual(bob['Tracking Token'], link.tracking_token)
        self.assertEqual(bob['Completed'], 'INCOMPLETE')
        self.assertEqual(bob['Reminder Eligible'], 'ELIGIBLE')
        self.assertEqual(bob['Reminder Status'], 'Sent (1)')
        alice = by_email['alice@example.com']
        self.assertEqual(alice['Completed'], 'COMPLETED')
        self.assertIsInstance(alice['Completed At'], datetime)
        self.assertEqual(alice['ODK Submission ID'], 'uuid:sub-A')
        self.assertEqual(alice['Reminder Eligible'], 'SUPPRESSED')
        self.assertEqual(alice['Reminder Status'], 'Suppressed')
        carol = by_email['carol@example.com']
        self.assertEqual(carol['Reminder Status'], 'Sent (1)')

        _h, clicks = self._table_rows(wb['Click Activity'])
        self.assertEqual(len(clicks), 2)
        bob_click = [e for e in clicks if e['Email'] == 'bob@example.com'][0]
        self.assertEqual(bob_click['Contact'], 'Bob Engaged')
        self.assertIsInstance(bob_click['Clicked At'], datetime)
        self.assertEqual(bob_click['Click Type'], 'Human')
        self.assertNotEqual(bob_click['Browser'], 'Unknown')

        _h, reminders = self._table_rows(wb['Reminder Activity'])
        sent = {r['Email']: r for r in reminders if r['Status'] == 'Sent'}
        self.assertEqual(set(sent), {'bob@example.com', 'carol@example.com'})
        self.assertEqual(sent['bob@example.com']['Reminder Sequence'], 1)
        self.assertIsInstance(sent['bob@example.com']['Sent At'], datetime)

        # Status cells: readable text AND distinct CF fills (never color alone).
        ws = wb['Recipient Lifecycle']
        treatments = {}
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                if rule.type == 'cellIs' and rule.formula:
                    fill = rule.dxf.fill
                    treatments.setdefault(rule.formula[0], set()).add(
                        fill.fgColor.rgb if fill is not None else None)
        self.assertEqual(set(treatments),
                         {'"COMPLETED"', '"INCOMPLETE"', '"ELIGIBLE"',
                          '"SUPPRESSED"', '"YES"'})
        fills = {text: next(iter(found))
                 for text, found in treatments.items()}
        self.assertEqual(len({fills['"COMPLETED"'], fills['"INCOMPLETE"'],
                              fills['"ELIGIBLE"'], fills['"SUPPRESSED"']}), 4)
        self.assertIsNone(fills['"YES"'])  # bold font only, no fill

    # 19. Summary totals match the report API ----------------------------------------------
    def test_19_summary_matches_report_api(self):
        api = get_campaign_full_report(self.v1)
        resp = self._export([self.v1.id])
        wb = self._wb(resp)
        kpis = self._kpis(wb['Summary'])
        self.assertEqual(kpis['Total Recipients'], api['kpi']['initial_eligible'])
        self.assertEqual(kpis['Emails Sent'], api['kpi']['sent_count'])
        self.assertEqual(kpis['Delivered'], api['kpi']['delivered_count'])
        self.assertEqual(kpis['Opened'], api['kpi']['unique_opened_count'])
        self.assertEqual(kpis['Clicked'], api['kpi']['unique_clicks_count'])
        self.assertEqual(kpis['Completed'], api['survey_conversion']['used_count'])
        self.assertEqual(kpis['Incomplete / UNUSED'], api['survey_conversion']['unused_count'])
        rates = self._rates(wb['Summary'])
        self.assertAlmostEqual(rates['Open Rate'],
                               api['kpi']['open_rate'] / 100)
        self.assertAlmostEqual(rates['Click Rate'],
                               api['kpi']['click_rate'] / 100)
        self.assertAlmostEqual(rates['Completion Rate'],
                               api['survey_conversion']['completion_rate'] / 100)
        self.assertAlmostEqual(
            rates['Delivery Rate'],
            kpis['Delivered'] / kpis['Emails Sent'])
        self.assertEqual(kpis['Reminders Sent'], 2)
        self.assertEqual(kpis['Reminder Eligible'], 2)
        self.assertEqual(kpis['Reminder Suppressed'], 3)
        table, combined = self._campaign_table(wb['Summary'])
        self.assertEqual(len(table), 1)
        self.assertEqual(table[0]['Recipients'], api['kpi']['initial_eligible'])
        self.assertEqual(table[0]['Sent'], api['kpi']['sent_count'])
        self.assertEqual(combined['Recipients'], api['kpi']['initial_eligible'])

    # 20. Existing CSV/XLSX/JSON behavior intact ----------------------------------------------
    def test_20_existing_exports_unchanged(self):
        base = '/api/campaigns/%d/report/link-recipients/' % self.v1.id
        resp = self.client.get(base)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.json()['total_count'], 5)
        resp = self.client.get(base, {'page': 2, 'page_size': 200})
        self.assertEqual(resp.json()['recipients'], [])

        resp = self.client.get(base, {'export': 'csv'})
        self.assertEqual(resp.status_code, 200)
        lines = resp.content.decode('utf-8').strip().splitlines()
        self.assertTrue(lines[0].startswith('Name,Email,Phone'))
        self.assertEqual(len(lines), 6)  # header + 5 recipients

        resp = self.client.get(base, {'export': 'xlsx'})
        self.assertEqual(resp.status_code, 200)
        simple = load_workbook(BytesIO(resp.content))
        self.assertEqual(simple.active.max_row, 6)

        # Same filters => CSV rows and workbook rows agree exactly.
        csv_emails = {line.split(',')[1] for line in lines[1:]}
        resp = self._export([self.v1.id])
        wb = self._wb(resp)
        _h, rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual({r['Email'] for r in rows}, csv_emails)

        clicks = self.client.get(
            '/api/campaigns/%d/report/link-clicks/' % self.v1.id).json()
        self.assertEqual(clicks['total_events'], 2)
        self.assertEqual(len(clicks['events']), 2)
        self.assertEqual(clicks['events'][0]['contact_email'], 'bob@example.com')

    # 21. Permissions + validation ----------------------------------------------------------------
    def test_21_permissions_and_validation(self):
        resp = self._export([self.v1.id], client=self.anon)
        self.assertIn(resp.status_code, (401, 403))
        self.assertEqual(self.client.get(URL).status_code, 405)
        resp = self._export([])
        self.assertEqual(resp.status_code, 400)
        resp = self._export([999999])
        self.assertEqual(resp.status_code, 400)
        resp = self._export([self.v1.id, 999998])
        self.assertEqual(resp.status_code, 400)
        resp = self.client.post(
            URL, data=json.dumps({'campaign_ids': 'not-a-list'}),
            content_type='application/json')
        self.assertEqual(resp.status_code, 400)

    # 23. Report page carries the Export Report dialog -------------------------------
    def test_23_report_page_shows_export_dialog(self):
        resp = self.client.get('/campaigns/%d/report/' % self.v1.id)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Export Campaign Report')
        self.assertContains(resp, 'id="export-modal"')
        self.assertContains(resp, 'id="export-campaign-card"')
        self.assertContains(resp, 'Pulse V1')
        self.assertContains(resp, 'onclick="openExportDialog()"')
        self.assertContains(resp, 'onclick="generateExportReport()"')
        self.assertContains(resp, 'id="export-lifecycle"')
        self.assertContains(resp, '/api/campaigns/report/xlsx/')

    # 25. Export dialog is fully wired (modal markup + every control) ----
    def test_25_export_dialog_wiring(self):
        # Regression test: the modal markup must render INSIDE a template
        # block (stray top-level markup in an {% extends %} child is
        # silently discarded by Django, which once left the button dead).
        resp = self.client.get('/campaigns/%d/report/' % self.v1.id)
        html = resp.content.decode('utf-8')
        for element_id in (
                'export-modal', 'export-error', 'export-campaign-card',
                'export-lifecycle', 'export-use-current', 'export-current-hint',
                'export-link', 'export-date-from',
                'export-date-to', 'export-csv-btn', 'export-csv-spinner',
                'export-csv-icon', 'export-sec-summary',
                'export-sec-campaign-info', 'export-sec-lifecycle',
                'export-sec-clicks', 'export-sec-incomplete',
                'export-sec-reminders', 'export-generate', 'export-spinner',
                'export-icon', 'export-generate-label'):
            self.assertIn('id="%s"' % element_id, html, element_id)
        # Single-campaign mode: no scope radios, no campaign picker.
        self.assertNotIn('export-scope', html)
        self.assertNotIn('export-campaign-picker', html)
        self.assertNotIn('exportCampaignsCache', html)
        self.assertNotIn('export-search', html)
        self.assertIn('const campaignId = %d;' % self.v1.id, html)
        # Every element the dialog JS looks up must exist exactly once.
        import re
        for used in set(re.findall(r"getElementById\('([^']+)'\)", html)):
            if used.startswith('export-') or used == 'lt-filter-link':
                self.assertIn('id="%s"' % used, html, used)

    # 24. Individual toolbar keeps Export Report + Export CSV -------------------
    def test_24_toolbar_exports(self):
        resp = self.client.get('/campaigns/%d/report/' % self.v1.id)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Export Report')
        self.assertContains(resp, 'id="export-csv-btn"')
        self.assertContains(resp, 'onclick="exportCampaignCsv()"')
        self.assertContains(resp, 'id="export-csv-spinner"')
        self.assertNotContains(resp, '/api/campaigns/%d/export/xlsx/' % self.v1.id)
        # Legacy endpoint itself stays live for backward compatibility.
        legacy = self.client.get('/api/campaigns/%d/export/xlsx/' % self.v1.id)
        self.assertEqual(legacy.status_code, 200)

    # 26. Campaign context is dynamic per page ---------------------------
    def test_26_campaign_context_dynamic(self):
        for campaign in (self.v1, self.v2):
            resp = self.client.get('/campaigns/%d/report/' % campaign.id)
            self.assertEqual(resp.status_code, 200)
            html = resp.content.decode('utf-8')
            self.assertIn('const campaignId = %d;' % campaign.id, html)
            self.assertIn(campaign.name, html)
            self.assertNotIn('export-scope', html)
            # The dialog itself offers no campaign selection (the sidebar
            # nav legitimately links to All Campaigns elsewhere).
            modal = html[html.index('id="export-modal"'):
                         html.index('id="export-generate-label"')]
            self.assertNotIn('All Campaigns', modal)
            self.assertNotIn('Select Campaign', modal)
            self.assertIn('This export covers this campaign only', modal)

    # 27. Individual export contains ONLY the current campaign -----------
    def test_27_individual_export_isolation(self):
        for campaign, other in ((self.v1, self.v2), (self.v2, self.v1)):
            resp = self._export([campaign.id])
            wb = self._wb(resp)
            _h, rows = self._table_rows(wb['Recipient Lifecycle'])
            self.assertTrue(rows)
            self.assertEqual({r['Campaign'] for r in rows}, {campaign.name})
            _h, clicks = self._table_rows(wb['Click Activity'])
            for e in clicks:
                if (e['Campaign'] or '').startswith('No records'):
                    continue
                self.assertEqual(e['Campaign'], campaign.name)
            table, _combined = self._campaign_table(wb['Summary'])
            self.assertEqual([r['Campaign'] for r in table], [campaign.name])
            self.assertNotIn(other.name, resp['Content-Disposition'])

    # 28. Individual CSV: headers, values, filename, encoding --------
    def test_28_campaign_csv(self):
        import csv as csv_module
        from io import StringIO
        resp = self.client.get('/api/campaigns/%d/export/csv/' % self.v1.id)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'text/csv; charset=utf-8')
        today = timezone.localdate().isoformat()
        self.assertIn('Campaign_Report_Pulse_V1_%s.csv' % today,
                      resp['Content-Disposition'])
        raw = resp.content.decode('utf-8-sig')
        self.assertTrue(resp.content.startswith(b'\xef\xbb\xbf'))  # BOM for Excel
        rows = list(csv_module.reader(StringIO(raw)))
        self.assertEqual(rows[0], ['Campaign', 'Name', 'Phone', 'Email',
                                   'Sent', 'Sent At', 'Opened', 'Opened At',
                                   'Clicked', 'Click Count',
                                   'First Clicked At', 'Tracking Token',
                                   'Completed', 'Completed At',
                                   'ODK Submission ID', 'Reminder Eligible',
                                   'Reminder Status', 'Reminder Reason'])
        self.assertEqual(len(rows), 6)  # header + 5 recipients
        self.assertTrue(all(r[0] == 'Pulse V1' for r in rows[1:]))
        by_email = {r[3]: r for r in rows[1:]}
        bob = by_email['bob@example.com']
        self.assertEqual(
            (bob[1], bob[4], bob[6], bob[8], bob[9], bob[12], bob[15],
             bob[16], bob[17]),
            ('Bob Engaged', 'YES', 'YES', 'YES', '1', 'INCOMPLETE',
             'ELIGIBLE', 'Sent (1)', 'Survey not completed'))
        import re
        stamp = re.compile(r'^\d{2} \w{3} \d{4}, \d{2}:\d{2} [AP]M$')
        self.assertRegex(bob[5], stamp)   # Sent At
        self.assertRegex(bob[10], stamp)  # First Clicked At
        alice = by_email['alice@example.com']
        self.assertEqual(alice[12], 'COMPLETED')
        self.assertRegex(alice[13], stamp)
        self.assertEqual(alice[14], 'uuid:sub-A')
        self.assertEqual((alice[15], alice[16], alice[17]),
                         ('SUPPRESSED', 'Suppressed', 'Survey completed'))
        dan = by_email['dan@example.com']
        self.assertEqual((dan[4], dan[5]), ('\u2014', '\u2014'))
        # Workbook lifecycle sheet and CSV agree on every recipient.
        wb = self._wb(self._export([self.v1.id]))
        _h, sheet_rows = self._table_rows(wb['Recipient Lifecycle'])
        self.assertEqual({r['Email'] for r in sheet_rows}, set(by_email))

    # 29. Individual CSV respects report filters ---------------------
    def test_29_campaign_csv_filters(self):
        import csv as csv_module
        from io import StringIO

        def emails(params):
            resp = self.client.get(
                '/api/campaigns/%d/export/csv/' % self.v1.id, params)
            self.assertEqual(resp.status_code, 200)
            rows = list(csv_module.reader(
                StringIO(resp.content.decode('utf-8-sig'))))
            return {r[3] for r in rows[1:]}, resp['Content-Disposition']

        got, _disp = emails({'status': 'clicked',
                             'contact_status': 'UNUSED'})
        self.assertEqual(got, {'bob@example.com'})
        got, disp = emails({'lifecycle': 'incomplete'})
        self.assertEqual(got, {'bob@example.com', 'carol@example.com',
                               'dan@example.com', 'eve@example.com'})
        self.assertIn('Incomplete_UNUSED', disp)
        got, _disp = emails({'lifecycle': 'completed'})
        self.assertEqual(got, {'alice@example.com'})

    # 30. CSV robustness: quoting + filename sanitization ------------
    def test_30_campaign_csv_robustness(self):
        import csv as csv_module
        from io import StringIO
        tricky = self._make_contact(
            'Oscar "Oz" O\u2019Brien, Jr.\nSecond Line',
            'oscar@example.com', '03009998888', 'VX-TRICKY', self.group1)
        CampaignMessage.objects.create(
            campaign=self.v1, contact=tricky, to_email='oscar@example.com',
            subject='s', message_type=CampaignMessage.MessageType.INITIAL,
            status=CampaignMessage.Status.DELIVERED,
            sent_at=timezone.now(), delivered_at=timezone.now())
        resp = self.client.get('/api/campaigns/%d/export/csv/' % self.v1.id)
        rows = list(csv_module.reader(
            StringIO(resp.content.decode('utf-8-sig'))))
        match = [r for r in rows[1:] if r[3] == 'oscar@example.com']
        self.assertEqual(len(match), 1)
        self.assertEqual(match[0][1], tricky.name)
        weird = Campaign.objects.create(
            name='Pulse: V1 / "Pilot" <test>?', subject='s', sender=self.sender,
            status=Campaign.Status.ACTIVE)
        resp = self.client.get('/api/campaigns/%d/export/csv/' % weird.id)
        self.assertEqual(resp.status_code, 200)
        disp = resp['Content-Disposition']
        bare = disp.split('filename=')[1].strip().strip('"')
        for ch in ':;"<>?|\\/':
            self.assertNotIn(ch, bare)
        self.assertIn('Campaign_Report_Pulse_V1_Pilot_test_', disp)

    # 31. Summary dashboard structure --------------------------------
    def test_31_summary_dashboard(self):
        resp = self._export([self.v1.id])
        wb = self._wb(resp)
        self.assertEqual(wb.active.title, 'Summary')
        ws = wb['Summary']
        self.assertEqual(ws['A1'].value, 'CAMPAIGN REPORT')
        self.assertEqual(
            ws['A2'].value,
            'Executive Performance & Recipient Lifecycle Overview')
        meta = self._meta(ws)
        self.assertEqual(meta['Campaign(s):'], 'Pulse V1')
        self.assertEqual(meta['Reporting period:'], 'All dates')
        self.assertIn('Applied filters:', meta)
        self.assertIn('Generated on:', meta)
        # KPI cards: large numeric values with thousands separators.
        value_cell = self._below(ws, 'Total Recipients')
        self.assertEqual(value_cell.value, 5)
        self.assertEqual(value_cell.number_format, '#,##0')
        self.assertEqual(value_cell.font.size, 20)
        # Rate strip: real fractions with percent formatting.
        for label in self.RATE_LABELS:
            cell = self._below(ws, label)
            self.assertIsInstance(cell.value, (int, float))
            self.assertGreaterEqual(cell.value, 0.0)
            self.assertLessEqual(cell.value, 1.0)
            self.assertEqual(cell.number_format, '0.0%')
        # Comparison is a real table; the total sits outside it, distinct.
        table = ws.tables['CampaignComparison']
        self.assertEqual(table.ref, 'A24:J25')
        self.assertEqual(table.tableStyleInfo.name, 'TableStyleMedium2')
        total_row = table.ref.split(':')[1]
        total_row = int(''.join(c for c in total_row if c.isdigit())) + 1
        self.assertEqual(ws.cell(row=total_row, column=1).value,
                         'Combined Total')
        self.assertEqual(ws.cell(row=total_row, column=2).fill.fgColor.rgb,
                         '000F172A')
        self.assertEqual(ws.cell(row=total_row, column=2).font.color.rgb,
                         '00FFFFFF')
        self.assertEqual(ws.freeze_panes, 'A7')
        self.assertFalse(ws.sheet_view.showGridLines)
        self.assertEqual(ws.sheet_properties.tabColor.rgb, '000F172A')
        self.assertTrue(ws.print_area.startswith("'Summary'!$A$1"))
        self.assertTrue(any(
            isinstance(c.value, str) and c.value.startswith('Definitions match')
            for row in ws.iter_rows() for c in row))

    # 32. Charts -----------------------------------------------------
    def test_32_charts_single_and_multi(self):
        for ids, cats in (([self.v1.id], "'Summary'!$A$25"),
                          ([self.v1.id, self.v2.id],
                           "'Summary'!$A$25:$A$26")):
            resp = self._export(ids)
            wb = self._wb(resp)
            ws = wb['Summary']
            self.assertEqual(
                self._chart_titles(ws),
                [('BarChart', 'Campaign Performance'),
                 ('DoughnutChart', 'Recipient Lifecycle')])
            charts = self._xlsx_xml(resp, 'xl/charts/')
            self.assertEqual(len(charts), 2)
            perf = next(x for x in charts.values()
                        if 'Campaign Performance' in x)
            life = next(x for x in charts.values()
                        if 'Recipient Lifecycle' in x)
            # Series read live comparison-table cells (titles + values).
            for ref in ("'Summary'!C24", "'Summary'!D24", "'Summary'!G24",
                        cats):
                self.assertIn(ref, perf)
            for color in ('64748B', '2563EB', '059669'):
                self.assertIn(color, perf)
            self.assertNotIn('3D', perf)
            table_ref = ws.tables['CampaignComparison'].ref
            anchor = int(''.join(
                c for c in table_ref.split(':')[1] if c.isdigit())) + 3
            kpis = self._kpis(ws)
            helper = {}
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value in ('Completed', 'Incomplete'):
                        helper[cell.value] = ws.cell(
                            row=cell.row, column=cell.column + 1).value
            self.assertEqual(helper['Completed'], kpis['Completed'])
            self.assertEqual(helper['Incomplete'], kpis['Incomplete / UNUSED'])
            self.assertIn(
                "'Summary'!$M$%d:$M$%d" % (anchor, anchor + 1), life)
            for color in ('059669', 'D97706'):
                self.assertIn(color, life)
            self.assertTrue(ws.column_dimensions['L'].hidden)
            self.assertTrue(ws.column_dimensions['M'].hidden)
            # Charts sit side by side without overlapping.
            import re
            drawing = self._xlsx_xml(resp, 'xl/drawings/drawing')['xl/drawings/drawing1.xml']
            anchors = re.findall(
                r'<from><col>(\d+)</col>.*?<row>(\d+)</row>', drawing, re.S)
            self.assertEqual(len(anchors), 2)
            self.assertEqual(anchors[0][1], anchors[1][1])  # same row band
            self.assertNotEqual(anchors[0][0], anchors[1][0])

    # 33. Navigation hyperlinks --------------------------------------
    def test_33_navigation_links(self):
        resp = self._export([self.v1.id])
        wb = self._wb(resp)
        links = {}
        for row in wb['Summary'].iter_rows():
            for cell in row:
                if (cell.hyperlink is not None and isinstance(cell.value, str)
                        and cell.value.startswith('Go to ')):
                    links[cell.value] = cell.hyperlink.location
        self.assertEqual(links, {
            'Go to Recipient Lifecycle': "'Recipient Lifecycle'!A2",
            'Go to Click Activity': "'Click Activity'!A2",
            'Go to Incomplete \u2014 Follow-up': "'Incomplete - UNUSED'!A2",
            'Go to Reminder Activity': "'Reminder Activity'!A2",
        })
        # Campaign names link to the detail sheet.
        name_cells = [cell for row in wb['Summary'].iter_rows() for cell in row
                      if cell.value == 'Pulse V1' and cell.column in (1, 2)]
        self.assertTrue(any(
            cell.hyperlink is not None
            and cell.hyperlink.location == "'Recipient Lifecycle'!A2"
            for cell in name_cells))
        for sheet in ('Recipient Lifecycle', 'Click Activity',
                      'Incomplete - UNUSED', 'Reminder Activity'):
            self.assertEqual(wb[sheet]['A1'].hyperlink.location,
                             "'Summary'!A1", sheet)
        # Disabled sections get no navigation buttons.
        resp = self._export(
            [self.v1.id],
            sections={'summary': True, 'campaign_info': False,
                      'lifecycle': True, 'clicks': False,
                      'incomplete': True, 'reminders': False})
        wb = self._wb(resp)
        labels = [cell.value for row in wb['Summary'].iter_rows()
                  for cell in row
                  if isinstance(cell.value, str)
                  and cell.value.startswith('Go to ')]
        self.assertEqual(sorted(labels),
                         ['Go to Incomplete \u2014 Follow-up',
                          'Go to Recipient Lifecycle'])
        # Summary off: detail sheets link back to Report Info.
        resp = self._export(
            [self.v1.id],
            sections={'summary': False, 'lifecycle': True, 'clicks': False,
                      'incomplete': False, 'reminders': False})
        wb = self._wb(resp)
        self.assertEqual(wb.active.title, 'Report Info')
        self.assertEqual(
            wb['Recipient Lifecycle']['A1'].hyperlink.location,
            "'Report Info'!A1")

    # 34. Number and date formatting -------------------------------
    def test_34_number_and_date_formats(self):
        resp = self._export([self.v1.id])
        wb = self._wb(resp)
        ws = wb['Summary']
        self.assertEqual(self._below(ws, 'Emails Sent').number_format,
                         '#,##0')
        rows, combined = self._campaign_table(ws)
        header_row = None
        for row in ws.iter_rows(values_only=False):
            if row[0].value == 'Campaign' and row[1].value == 'Recipients':
                header_row = row[0].row
        self.assertEqual(
            ws.cell(row=header_row + 1, column=2).number_format, '#,##0')
        wl = wb['Recipient Lifecycle']
        headers, data = self._table_rows(wl)
        sent_col = headers.index('Sent At') + 1
        for values in wl.iter_rows(min_row=3, values_only=False):
            if values[0].value is None:
                continue
            cell = values[sent_col - 1]
            if cell.value == '\u2014':
                continue
            self.assertIsInstance(cell.value, datetime)
            self.assertEqual(cell.number_format, 'DD MMM YYYY HH:MM')

    # 35. Empty selection: no charts, header-only tables ------------
    def test_35_empty_selection(self):
        resp = self._export([self.v1.id],
                            filters={'search': 'zzz-no-such-person'})
        wb = self._wb(resp)
        ws = wb['Summary']
        self.assertEqual(ws._charts, [])
        self.assertEqual(
            self._xlsx_xml(resp, 'xl/charts/'), {})
        for sheet, ref in (('Recipient Lifecycle', 'A2:R2'),
                           ('Click Activity', 'A2:K2'),
                           ('Incomplete - UNUSED', 'A2:M2'),
                           ('Reminder Activity', 'A2:G2')):
            table = next(iter(wb[sheet].tables.values()))
            self.assertEqual(table.ref, ref, sheet)
            self._assert_empty_note(wb[sheet])
        kpis = self._kpis(ws)
        for label in self.KPI_LABELS:
            self.assertEqual(kpis[label], 0, label)
        for label in self.RATE_LABELS:
            self.assertEqual(self._below(ws, label).value, 0.0, label)
        _rows, combined = self._campaign_table(ws)
        self.assertEqual(combined['Recipients'], 0)

    # 36. Print setup and workbook metadata ------------------------
    def test_36_print_and_metadata(self):
        resp = self._export([self.v1.id, self.v2.id])
        wb = self._wb(resp)
        self.assertEqual(wb.properties.creator, 'IRIS Marketing')
        self.assertEqual(wb.properties.lastModifiedBy, 'IRIS Marketing')
        self.assertIn('Campaign Report', wb.properties.title)
        self.assertIn('Pulse V1', wb.properties.title)
        self.assertIn('Filters:', wb.properties.description)
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            self.assertEqual(ws.page_setup.orientation, 'landscape', sheet)
            self.assertTrue(ws.sheet_properties.pageSetUpPr.fitToPage, sheet)
            self.assertEqual(ws.page_setup.fitToWidth, 1, sheet)
            self.assertEqual(ws.page_setup.fitToHeight, 0, sheet)
            self.assertEqual(ws.oddFooter.center.text, 'Page &P of &N', sheet)
            self.assertTrue(ws.print_area, sheet)

    # 22. Filenames -----------------------------------------------------------------------------------
    def test_22_filenames(self):
        today = timezone.localdate().isoformat()
        self.assertEqual(
            build_report_filename(['Pulse V1']),
            'Campaign_Report_Pulse_V1_%s.xlsx' % today)
        self.assertEqual(
            build_report_filename(['Pulse V1', 'Pulse V2']),
            'Campaign_Report_Pulse_V1_Pulse_V2_%s.xlsx' % today)
        self.assertEqual(
            build_report_filename(['Pulse V1'], lifecycle='incomplete'),
            'Campaign_Report_Pulse_V1_Incomplete_UNUSED_%s.xlsx' % today)
        self.assertEqual(
            build_report_filename(['Pulse V1 (Pilot) #2!'], search='bob'),
            'Campaign_Report_Pulse_V1_Pilot_2_Filtered_%s.xlsx' % today)
        self.assertEqual(
            build_report_filename([]),
            'Campaign_Report_Campaigns_%s.xlsx' % today)
        long_name = 'C' * 200
        name = build_report_filename([long_name])
        self.assertLessEqual(len(name), 130)
        self.assertTrue(name.endswith('.xlsx'))
        # Endpoint echoes the dynamic filename.
        resp = self._export([self.v1.id, self.v2.id],
                            filters={'lifecycle': 'incomplete'})
        self.assertIn('Campaign_Report_Pulse_V1_Pulse_V2_Incomplete_UNUSED_',
                      resp['Content-Disposition'])
